import sys,json
from pathlib import Path
sys.path.insert(0,str(Path.cwd()))
sys.path.append(str(Path('backend/.venv/Lib/site-packages').resolve()))
from pptx import Presentation
from docx import Document
from ai_server.app.idea_proposal import prepare_idea_report
from ai_server.app.proposal_calculation_basis import calculation_basis
from ai_server.app.proposal_presentation_v4 import _scenario_for_display
r=prepare_idea_report(json.loads(Path('storage/previews/jeju_revision_20260911/original.json').read_text(encoding='utf-8')))
s,_=_scenario_for_display(r)
b=calculation_basis(r)
ppt=Presentation('output/jeju_revision_20260911/제주시_관광전략기획안_인포그래픽_v2.pptx')
charts=[]
for slide in ppt.slides:
 for sh in slide.shapes:
  if sh.has_chart:
   charts.append({'series':[{'name':s.name,'values':list(s.values)} for s in sh.chart.series]})
doc=Document('output/jeju_revision_20260911/제주시_관광전략기획안_Word_최종.docx')
text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
print(json.dumps({'months':s['categories'],'totals':b['totals'],'estimate':b['estimate']['total_krw'],'images':len(doc.inline_shapes),'source_count':len(r['evidence_sources']),'ppt_charts':charts,'missing_source_ids':[x['source_id'] for x in r['evidence_sources'] if x.get('source_id') and x['source_id'] not in text]},ensure_ascii=True,indent=2))
