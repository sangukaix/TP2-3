"""Read-only package/data QA for the actual Wonju PPT; no model calls."""
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import sys

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from pptx import Presentation
from ai_server.app.report_projection import select_report_forecast
from ai_server.app.proposal_evidence import complete_source_records,build_reference_estimate
from ai_server.app.proposal_slide_content import _source_groups

OUT=Path(__file__).resolve().parent
report=json.loads((OUT/'enriched.json').read_text(encoding='utf-8'))
dest=ROOT/'storage/previews/원주시_관광전략기획안_실제데이터_수정본.pptx'
deck=Presentation(dest)
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
alltext='\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
notes='\n'.join(s.notes_slide.notes_text_frame.text for s in deck.slides)
appendix=''.join(''.join(s.text.split()) for slide in list(deck.slides)[10:] for s in slide.shapes if s.has_text_frame)
rows=select_report_forecast(report)['rows']
charts=[s.chart for s in deck.slides[4].shapes if s.has_chart]
assert len(charts)==2
for chart,key,unit in zip(charts,['visitors','spending_krw'],[1e4,1e8]):
    actual=chart.series[0].values
    assert len(actual)==len(rows)
    assert all(abs(v*unit-r[key])<.1 for v,r in zip(actual,rows))
for source in complete_source_records(report):
    assert source['source_id'] in notes,source['source_id']
for title,items in _source_groups(report):
    for text,url in items:
        assert ''.join(text.split()) in appendix,(title,text)
assert '운영 후 실제 확인' not in alltext
assert 'AI 기획·검수' not in alltext
assert 'CORE OPERATING MODEL · SAMPLE' not in alltext
assert '월 방문자 수\n-2.7%' not in alltext
assert '감사합니다' in '\n'.join(s.text for s in deck.slides[-1].shapes if s.has_text_frame)
estimate=build_reference_estimate(report)
assert estimate['total_krw']==sum(i['amount'] for i in estimate['items'])==82_500_000
empty=[]
with ZipFile(dest) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.startswith('ppt/slides/slide') and name.endswith('.xml'):
            root=ET.fromstring(z.read(name))
            for shape in root.findall('.//p:sp',ns):
                if shape.find('.//p:ph',ns) is not None and not ''.join(shape.itertext()).strip():
                    empty.append(name)
assert not empty,empty
with ZipFile(ROOT/'ai_server/app/templates/tourism_strategy_12_slide_template.pptx') as old, ZipFile(ROOT/'ai_server/app/templates/tourism_strategy_12_slide_template_v6.pptx') as new:
    themes=[n for n in old.namelist() if n.startswith('ppt/theme/') and n.endswith('.xml')]
    assert all(old.read(n)==new.read(n) for n in themes)
    assert all(old.read(f'ppt/slides/slide{i}.xml')==new.read(f'ppt/slides/slide{i}.xml') for i in [1,2,3,4,6,7,10,12])
print(json.dumps({'slides':len(deck.slides),'sources_preserved':len(complete_source_records(report)),
  'monthly_chart_rows':len(rows),'models':len(report['ml_analysis']['evaluation']['metrics']),
  'estimate_krw':estimate['total_krw'],'all_source_names_visible':True,
  'original_themes_preserved':True,'untouched_template_slides':8,'empty_placeholders':empty},ensure_ascii=False))
