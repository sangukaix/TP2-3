"""Approved editable charts/tables; facts remain separate from planning targets."""
from .proposal_layout_v7 import *
from . import proposal_layout_v7 as v7
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_DATA_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.oxml.xmlchemy import OxmlElement


def centered(shape):
    for p in shape.text_frame.paragraphs:p.alignment=PP_ALIGN.CENTER
    return shape


def compact(value, key):
    unit=1e4 if key=='visitors' else 1e8
    suffix='만 명' if key=='visitors' else '억 원'
    return f'{value/unit:,.0f}{suffix}'


def comparison_chart(slide,name,x,y,w,h,scenario,key,divisor):
    sh=v7.chart(slide,name,x,y,w,h,scenario,key,divisor);c=sh.chart
    # Rename both XML cache and embedded workbook consistently by replacing data.
    data=CategoryChartData();data.categories=scenario['categories'][:3]
    data.add_series('ML 기준 전망',[round(v/divisor,8) for v in scenario['baseline_'+key][:3]])
    if scenario['has_target']:data.add_series('목표 KPI',[round(v/divisor,8) for v in scenario['target_'+key][:3]])
    c.replace_data(data)
    values=[v/divisor for v in scenario['baseline_'+key][:3]]
    if scenario['has_target']:values += [v/divisor for v in scenario['target_'+key][:3]]
    spread=max(max(values)-min(values),abs(max(values))*.12,1)
    raw=spread/4;scale=10**math.floor(math.log10(raw))
    step=next(n*scale for n in (1,2,5,10) if n*scale>=raw)
    c.value_axis.minimum_scale=max(0,math.floor((min(values)-spread*.35)/step)*step)
    c.value_axis.maximum_scale=math.ceil((max(values)+spread*.3)/step)*step
    c.value_axis.major_unit=step
    c.has_legend=False
    c.category_axis.tick_label_position=XL_TICK_LABEL_POSITION.LOW
    c.plots[0].has_data_labels=True
    for series,color,position in zip(c.series,(BLUE,ORANGE),(XL_DATA_LABEL_POSITION.BELOW,XL_DATA_LABEL_POSITION.ABOVE)):
        labels=series.data_labels
        labels.position=position
        labels.number_format='0"'+('만' if key=='visitors' else '억')+'"'
        labels.number_format_is_linked=False
        labels.show_value=True
        labels.show_legend_key=False
        labels.show_category_name=False
        labels.show_series_name=False
        labels.font.name='맑은 고딕';labels.font.size=Pt(12);labels.font.bold=True
        labels.font.color.rgb=RGBColor.from_string(color)
    # Let PowerPoint reserve axis/label margins inside a shorter chart canvas.
    # Manual plot coordinates are interpreted differently by desktop renderers.
    sh.height=round((h-46)*EMU)
    plot=c._chartSpace.chart.plotArea
    for old in plot.findall('{http://schemas.openxmlformats.org/drawingml/2006/chart}layout'):plot.remove(old)
    # Reserve a separate bottom row so PowerPoint cannot place the legend over months.
    for i,(label,color) in enumerate([('ML 기준 전망',BLUE)]+([('목표 KPI',ORANGE)] if scenario['has_target'] else [])):
        lx=x+164+i*181
        rect(slide,f'{name}-legend-line-{i}',lx,y+h-16,25,3,color)
        text(slide,f'{name}-legend-label-{i}',label,lx+30,y+h-28,155,30,18,SLATE)
    return sh


def detail(prs,report):
    from .proposal_presentation_v4 import _selected_candidate,_scenario_for_display
    s=prs.slides[3];strategy=base._first_strategy(report);candidate=_selected_candidate(report)
    scenario,demo=_scenario_for_display(report)
    header(s,'사업 설계와 3개월 목표')
    target='목표 방문자 수·관광소비액은 목표율 확정 후 표시합니다.'
    if scenario and scenario['has_target']:
        target=f"3개월 월별 합계\n방문자 {compact(sum(scenario['target_visitors'][:3]),'visitors')}\n관광소비액 {compact(sum(scenario['target_spending'][:3]),'spending')}"
    blocks=[('사업내용',strategy.get('solution') or candidate.get('mechanism') or '사업내용 확인 필요'),
            ('사업기간',f"{strategy.get('timeframe') or '일정 협의'}\n준비·협의 → 시범 운영 → 성과 확인"),('목표 KPI',target)]
    bodies=[]
    for i,(label,body) in enumerate(blocks):
        x=106+i*470;rect(s,f'block-{i}',x,216,448,207,WHITE)
        rect(s,f'block-rail-{i}',x,216,6,207,CYAN if i==1 else BLUE)
        text(s,f'block-title-{i}',label,x+18,231,414,40,25,BLUE,True)
        bodies.append(text(s,f'block-body-{i}',str(body),x+18,281,414,132,23,SLATE))
    size=min(r.font.size.pt/.75 for sh in bodies for p in sh.text_frame.paragraphs for r in p.runs)
    for sh,(_,body) in zip(bodies,blocks):fit_text(sh,str(body),size,SLATE)
    for i,(key,label,divisor) in enumerate([('visitors','관광객 방문자 수',1e4),('spending','관광소비액',1e8)]):
        x=106+i*714;rect(s,f'chart-panel-{i}',x,446,674,350,WHITE)
        centered(text(s,f'chart-title-{i}',label,x+20,456,634,43,28,BLUE,True))
        if scenario:comparison_chart(s,f'comparison-chart-{i}',x+12,510,650,277,scenario,key,divisor)
        else:centered(text(s,f'chart-missing-{i}','사업 기간의 예측치 미제공',x,585,674,90,25,SLATE))
    period=str((report.get('ml_analysis') or {}).get('source_period') or '저장 학습기간')
    note=f'ML 기준: {period} 월별 지표의 1·3·12개월 전 값과 계절 정보를 저장 모델에 입력. 목표: 기준 전망 × (1 + 최종월 목표율 × 경과월/전체월).\n목표율은 사례 실적을 참고한 계획 가정이며, 구체적인 채택 비율과 출처는 목표 설정 근거에 제시합니다.'
    text(s,'comparison-note',note,106,811,1388,63,20,SLATE)


def diverse_cases(report):
    from .proposal_presentation_v4 import _selected_case,_selected_candidate
    from .agents.case_study_agent import _case_mechanism_family
    selected,assessment=_selected_case(report)
    chosen=[*(_selected_candidate(report).get('case_source_ids') or []),*((report.get('planning_decision') or {}).get('recommended_case_ids') or [])]
    if selected.get('source_id') not in chosen:selected={};assessment={}
    pool=[r for r in complete_source_records(report) if r.get('source_type') in ('benchmark_case','case_study','official_case')]
    result=[selected] if selected else [];seen={_case_mechanism_family(selected)} if selected else set()
    for unique in (True,False):
        for r in pool:
            if len(result)>=4:break
            if any(r.get('source_id')==v.get('source_id') for v in result):continue
            family=_case_mechanism_family(r)
            if unique and family in seen:continue
            result.append(r);seen.add(family)
    return result,selected,assessment


def cases(prs,report):
    s=prs.slides[5];header(s,'지역별 참고 사례');rows,selected,assessment=diverse_cases(report)
    for i in range(4):
        x=106+(i%2)*714;y=223+(i//2)*313;r=rows[i] if i<len(rows) else {}
        rect(s,f'case-panel-{i}',x,y,674,290,PALE if i==0 and selected else WHITE)
        text(s,f'case-number-{i}',f'{i+1:02d}',x+20,y+16,65,50,33,CYAN,True)
        text(s,f'case-badge-{i}','선정 사례' if i==0 and selected else '기타 유사 사례',x+98,y+20,545,36,20,BLUE,True)
        parts=copy_text(r.get('summary')).split(' · ')
        label=' · '.join(parts[:2]) if len(parts)>=3 else r.get('title')
        # 줄 계산은 실제 상자보다 좁게 하여 PowerPoint 폰트 대체에도 여백을 유지합니다.
        title=text(s,f'case-title-{i}',label or '추가 공식 사례 미확보',x+28,y+73,534,77,27,INK,True)
        title.width=594*EMU
        title.text_frame.word_wrap=True
        body=r.get('operating_model') or r.get('summary') or '확인된 공식 사례가 없어 임의로 추가하지 않습니다.'
        if i==0 and assessment.get('adaptation'):body='적용 조건: '+assessment['adaptation']
        description=text(s,f'case-body-{i}',copy_text(body),x+28,y+159,534,113,21,SLATE)
        description.width=594*EMU
        description.text_frame.word_wrap=True


def table(slide,name,rows,x,y,w,h,widths,size=22):
    sh=native_table(slide,name,rows,x,y,w,h,widths,size)
    for row in sh.table.rows:
        for cell in row.cells:cell.vertical_anchor=MSO_ANCHOR.MIDDLE
    return sh


def roadmap(prs,report):
    s=prs.slides[6];header(s,'4단계 실행 가이드')
    steps=[v for v in base._first_strategy(report).get('implementation_steps') or [] if isinstance(v,dict)]
    groups=[[v] for v in steps[:3]]+[steps[3:]]
    while len(groups)<4:groups.append([])
    roles=('사업 총괄','운영·정산 담당','현장 운영 담당','사업·평가 담당')
    labels=('준비와 협의','운영 절차 점검','시범 운영','실적 확인과 평가')
    guidance=('참여 업체를 모집하고 지원 대상·신청 방법·지원 한도를 정합니다.',
              '신청부터 증빙 확인·지급까지 시험합니다. 취소·중복 신청 처리 방법도 정합니다.',
              '작은 범위에서 먼저 운영합니다. 이용 안내와 안전을 점검하고 일별 신청·이용 기록을 남깁니다.',
              '운영 실적과 비용·민원을 매주 확인합니다. 종료 후 기준기간과 비교해 결과를 정리합니다.')
    rows=[['단계·기간','담당 역할(제안)','세부 실행 업무']]
    for i,group in enumerate(groups):
        x=106+i*355
        rect(s,f'pipeline-{i}',x,234,316,73,BLUE if i%2==0 else CYAN)
        centered(text(s,f'step-{i}-label',f'{i+1:02d}  {labels[i]}',x+4,250,308,46,24,WHITE,True))
        if i<3:text(s,f'arrow-{i}','→',x+320,251,37,42,26,CYAN,True)
        schedule=' / '.join(dict.fromkeys(copy_text(v.get('schedule')) for v in group if v.get('schedule')))
        centered(text(s,f'step-{i}-period',schedule or '일정 협의',x,318,316,61,21,SLATE))
        saved=' / '.join(copy_text(v.get('task')) for v in group)
        rows.append([f'{i+1:02d}\n{schedule or "일정 협의"}',roles[i],(saved+'\n' if saved else '')+guidance[i]])
    table(s,'execution-table',rows,106,401,1388,414,[225,230,933],22)
    text(s,'roadmap-note','담당 역할과 업무는 실행 제안입니다. 착수 전에 담당자·협력처와 적용 조건을 확정합니다.',106,838,1388,42,19,SLATE)


def kpi(prs,report):
    from .proposal_presentation_v4 import _scenario_for_display
    from .proposal_slide_content import _people,_money
    s=prs.slides[7];scenario,demo=_scenario_for_display(report)
    header(s,'머신러닝 예측값과 목표 KPI')
    for k,(key,label,fmt) in enumerate([('visitors','관광객 방문자 수',_people),('spending','관광소비액',_money)]):
        y=215+k*280
        text(s,f'kpi-title-{k}',label,106,y,850,39,26,BLUE,True)
        rows=[['기준월','머신러닝 예측값','목표 KPI','기준 대비 증가']]
        bs=scenario['baseline_'+key][:3] if scenario else []
        ts=scenario['target_'+key][:3] if scenario and scenario['has_target'] else []
        for i,m in enumerate(scenario['categories'][:3] if scenario else ['미제공']):
            b=bs[i] if bs else None;t=ts[i] if ts else None
            rows.append([m,fmt(b) if b is not None else '미제공',fmt(t) if t is not None else '목표율 미입력',f'↑ {fmt(t-b)} · {(t/b-1)*100:.2f}%' if t is not None and b else '산출 보류'])
        rows.append(['3개월 월별 합계',fmt(sum(bs)) if bs else '미제공',fmt(sum(ts)) if ts else '목표율 미입력',f'↑ {fmt(sum(ts)-sum(bs))} · {(sum(ts)/sum(bs)-1)*100:.2f}%' if ts and sum(bs) else '산출 보류'])
        table(s,f'kpi-{key}-table',rows,106,y+49,1388,212,[225,350,350,463],21)
        if ts:
            centered(text(s,f'kpi-total-{k}',f'3개월 목표 추가 규모  ↑ {fmt(sum(ts)-sum(bs))}',918,y,576,43,23,ORANGE,True))
    note='방문자 합계는 월별 합계로 월 간 중복을 포함합니다. 목표 KPI는 예측값에 입력 목표율을 단계 적용한 계획값이며, 보장된 사업 효과가 아닙니다.'
    if scenario and scenario.get('spending_calculation') == 'visitor_linked':
        ratios=' · '.join(f'{month} {ratio:,.0f}원' for month,ratio in zip(
            scenario['categories'][:3],scenario['spending_per_visit_proxy'][:3]))
        note=(f'월별 소비/방문 비율 = ML 소비액 ÷ ML 방문 수  |  {ratios}\n'
              '목표 소비액 = ML 소비액 + (목표 방문 수 − ML 방문 수) × 해당 월 소비/방문 비율\n'
              '월별 증가율은 방문·소비가 같습니다. 합계 증가율 = (목표 합계 ÷ ML 합계 − 1) × 100; 월별 가중치에 따라 달라집니다.\n'
              '비율은 기획용 가정이며 실제 객단가·보장 효과가 아닙니다. 방문 합계는 월 간 중복을 포함합니다.')
    elif scenario and scenario.get('has_target'):
        note=('월별 목표 = 해당 월 ML 전망 × (1 + 최종월 목표율 × 경과 월 수 ÷ 전체 월 수). 방문·소비 목표율은 각각 적용합니다.\n'
              '합계 증가율 = (목표 합계 ÷ ML 합계 − 1) × 100. 방문 합계는 월 간 중복을 포함하며, 보장된 사업 효과가 아닙니다.')
    if demo or report.get('layout_target_sample'):note+=' 이번 파일의 목표율은 디자인 확인용 예시입니다.'
    text(s,'kpi-calculation',note,106,780,1388,112,18,SLATE)


def budget(prs,report):
    from .proposal_slide_content import _money
    from .proposal_evidence import saved_budget_view
    s=prs.slides[8];header(s,'견적');e=build_reference_estimate(report)
    saved = None if report.get('reference_estimate') else (saved_budget_view(report) if report.get('generation_mode') != 'offline_sample' else None)
    e = report.get('reference_estimate') or e
    text(s,'estimate-label','시범 운영 참고 금액',106,228,790,48,28,SLATE,True)
    text(s,'estimate-total',saved['total_label'] if saved else f"총 {_money(e['total_krw'])}",950,218,544,65,42,BLUE,True)
    rows=[['항목','산출 기준(가정)','금액']]+[[r['name'],r['basis'],f"{r['amount']/10000:,.0f}만원"] for r in e['items']]
    if saved:
        rows = [['항목', '저장 기획안의 산출 근거', '구분']] + saved['rows']
    sh=table(s,'estimate-table',rows,130,330,1340,376,[345,670,325],22)
    for i in range(len(rows)):
        for p in sh.table.cell(i,2).text_frame.paragraphs:p.alignment=PP_ALIGN.RIGHT
    rect(s,'estimate-total-rule',130,720,1340,3,CYAN)
    text(s,'estimate-note',saved['note'] if saved else '수량과 단가를 가정한 임시 참고 견적이며 확정 견적이 아닙니다. 실제 계약 범위·세금·단가 확인 후 조정합니다.',130,761,1340,74,23,SLATE)


def provenance(prs,report):
    pages=v7.provenance(prs,report)
    metrics=((report.get('ml_analysis') or {}).get('evaluation') or {}).get('metrics') or {}
    weak=[name for key,name in [('visitors','방문자'),('spending_krw','소비액')]
          if (metrics.get(key) or {}).get('model_reliability')=='below_baseline_on_test']
    warning=('참고 전망: '+ '·'.join(weak)+' 모델의 검증 오차가 전년동월 기준선보다 큽니다.') if weak else ''
    if select_report_forecast(report).get('notes'):warning+=' 예측거리·누락 범위의 해석 주의가 필요합니다.'
    if warning:
        for s in prs.slides:
            if any(getattr(sh,'has_text_frame',False) and sh.name=='title' and sh.text=='근거·데이터·출처' for sh in s.shapes):
                sh=next(sh for sh in s.shapes if sh.name=='subtitle')
                sh.height=65*EMU;fit_text(sh,sh.text+'\n'+warning,19,SLATE)
    return pages


def finish(prs,report):
    v7.finish(prs,report)
    sid=prs.slides._sldIdLst[8];prs.part.drop_rel(sid.rId);prs.slides._sldIdLst.remove(sid)
    labels=['프로젝트 개요','사업 설계와 3개월 목표','유사지역사례','4단계 실행 가이드','머신러닝 예측값과 목표 KPI','견적','근거·데이터·출처','맺음말']
    toc=prs.slides[1]
    for i in range(10):
        col,row=divmod(i,5)
        for sh in list(toc.shapes):
            if sh.name.startswith(f'toc-{col}-{row}'):
                if i>=len(labels):sh._element.getparent().remove(sh._element)
                elif sh.name.endswith('-label'):fit_text(sh,labels[i],26,INK,True)
    for sh in list(prs.slides[-1].shapes):
        if getattr(sh,'has_text_frame',False) and (sh.name=='org-owner' or any(t in sh.text.upper() for t in ('FINAL','TOUR INSIGHT','관광 전략 검토용','관광전략 검토용'))):
            sh._element.getparent().remove(sh._element)
    assert len(prs.slides)>=10
