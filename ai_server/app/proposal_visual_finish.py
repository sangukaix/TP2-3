"""Final common visual hierarchy. Preserve report text, numbers and source links."""
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt
from .proposal_layout_v7 import EMU, BLUE, CYAN, SLATE, WHITE, text, rect, fit_text


def restyle(shape, size, color=BLUE, bold=True):
    if not shape.has_text_frame:
        return
    value=shape.text
    # fit_text recomputes wrapping from the complete original string.
    fit_text(shape,value,size,color,bold)


def finish_design(prs, report):
    count=len(prs.slides)
    for index,slide in enumerate(prs.slides):
        title=next((s for s in slide.shapes if s.name=='title' and s.has_text_frame),None)
        if index>=2 and index<count-1 and title:
            title.top=86*EMU;title.height=77*EMU
            restyle(title,42)
            rule=next((s for s in slide.shapes if s.name=='title-rule'),None)
            if rule:
                rule.height=1*EMU;rule.fill.fore_color.rgb=RGBColor.from_string('D4DFE8')
            header_x = 561 if index == 2 and any(s.name == 'visitor-photo' for s in slide.shapes) else 106
            rect(slide,'editorial-accent',header_x,179,79,4,CYAN)
            text(slide,'editorial-kicker','관광 전략기획안',header_x,45,800 if header_x == 561 else 1100,28,15,SLATE)
            page=text(slide,'editorial-page',f'{index+1:02d} / {count:02d}',1393,45,101,28,15,SLATE)
            page.text_frame.paragraphs[0].alignment=PP_ALIGN.RIGHT
        for shape in list(slide.shapes):
            name=shape.name
            if name.startswith(('execution-panel-','case-panel-')):
                # Clear white planes and thin cyan rails echo the supplied example.
                rect(slide,name+'-accent',shape.left/EMU,shape.top/EMU,3,shape.height/EMU,CYAN)
            if name.startswith('s3-goal-panel-'):
                shape.fill.fore_color.rgb=RGBColor.from_string(WHITE)
                rect(slide,name+'-accent',shape.left/EMU,shape.top/EMU,3,shape.height/EMU,CYAN)
            if name=='s3-thesis':restyle(shape,30,'1D2227')
            if name=='s3-detail-label':restyle(shape,24)
            if name=='selection-case':restyle(shape,27)
            if name.startswith('case-badge-'):restyle(shape,27)
            if name.startswith('case-title-'):restyle(shape,24,'1D2227')
            if name in ('kpi-intro','pipeline-intro'):restyle(shape,24,SLATE,False)
            if name=='kpi-reason-title':restyle(shape,29)
            if name=='kpi-reason':restyle(shape,26,SLATE,False)
            if name=='kpi-decision':restyle(shape,30)
            if name=='kpi-step-explanation':restyle(shape,22,SLATE,False)
            if name.startswith('kpi-impact-title-'):restyle(shape,25)
            if name.startswith('kpi-impact-value-'):restyle(shape,41,'FF6B00')
            if name.startswith('kpi-impact-period-'):restyle(shape,20,SLATE,False)
            if name=='result-title':restyle(shape,30)
            if name=='result-change':restyle(shape,38,'FF6B00')
            if name.startswith('block-title-'):restyle(shape,23)
            if name.startswith('chart-title-'):
                restyle(shape,26)
                for p in shape.text_frame.paragraphs:p.alignment=PP_ALIGN.CENTER
            if name=='estimate-total':restyle(shape,35)
            if shape.has_text_frame:
                # Keep clickable source metadata, but avoid a page of underlines.
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if run.hyperlink.address:
                            run.font.underline=False
                            run.font.color.rgb=RGBColor.from_string('39718B')
        # All native tables share deliberate compact paragraph spacing.
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            p.space_before=p.space_after=Pt(0)
