"""승인된 12장 편집형 템플릿에 검증된 기획안 데이터를 채웁니다.

이 모듈은 분석·예측을 새로 수행하지 않습니다. 이미 검증된 ReportResponse의
관측값, 저장 ML 자연추세, 사용자 목표, 5단계 실행안만 템플릿에 배치합니다.
"""

from __future__ import annotations

import re
from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlparse

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN

from .presentation_theme import download_images, select_image_sources
from .report_projection import execution_target, select_report_forecast, target_series, visitor_linked_spending
from .report_review_status import review_label


# 시각 디자인은 코드로 다시 그리지 않고 승인된 편집형 PPTX를 그대로 사용합니다.
PRESENTATION_TEMPLATE_PATH = (
    Path(__file__).resolve().parent / 'templates' / 'tourism_strategy_12_slide_template_legacy_v1.pptx'
)
PRESENTATION_RENDER_VERSION = 'pptx-12-template-v2-period'

NAVY = RGBColor(0x00, 0x4E, 0xA2)
CYAN = RGBColor(0x00, 0xC1, 0xD4)
ORANGE = RGBColor(0xFF, 0x6B, 0x00)


def _compact(value: Any, *, limit: int = 180) -> str:
    """긴 문장을 템플릿의 글상자 크기에 맞게 공백 정리 후 줄입니다."""
    text = re.sub(r'\s+', ' ', str(value or '')).strip()
    if len(text) <= limit:
        return text
    shortened = text[:limit].rsplit(' ', 1)[0].rstrip(',. ')
    return f'{shortened or text[:limit]}…'


def _sentences(value: Any, *, limit: int = 4, length: int = 105) -> list[str]:
    """서술형 입력을 슬라이드에 들어갈 짧은 문장 단위로 나눕니다."""
    text = re.sub(r'\s+', ' ', str(value or '')).strip()
    if not text:
        return []
    parts = [
        item.strip(' ·-')
        for item in re.split(r'(?<=[.!?다요])\s+|[;•]\s*', text)
        if item.strip()
    ]
    if len(parts) == 1 and len(parts[0]) > length:
        parts = [item.strip() for item in re.split(r',\s*', parts[0]) if item.strip()]
    return [_compact(item, limit=length) for item in parts[:limit]]


def _as_float(value: Any) -> float:
    """차트 입력값의 None·문자열을 안전한 실수로 바꿉니다."""
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _month_label(value: Any) -> str:
    """202609 또는 2026.09를 26.09 형식으로 표시합니다."""
    digits = re.sub(r'\D', '', str(value or ''))
    if len(digits) >= 6:
        return f'{digits[2:4]}.{digits[4:6]}'
    return str(value or '')


def _format_people(value: float) -> str:
    """방문자 수를 만 명 단위로 표시합니다."""
    return f'{value / 10_000:,.0f}만 명'


def _format_money(value: float) -> str:
    """관광소비액을 억 원 또는 조 원 단위로 표시합니다."""
    if abs(value) >= 1_000_000_000_000:
        return f'{value / 1_000_000_000_000:,.3f}조 원'
    return f'{value / 100_000_000:,.0f}억 원'


def _metric_card_value(metric: str, value: str) -> str:
    """긴 원 단위·명 단위 숫자를 카드에서 읽기 쉬운 축약 단위로 바꿉니다."""
    match = re.search(r'-?[\d,]+(?:\.\d+)?', str(value or ''))
    if not match:
        return _compact(value or '자료 없음', limit=18)
    number = _as_float(match.group(0).replace(',', ''))
    if '소비' in metric and '비율' not in metric:
        return _format_money(number)
    if '방문자' in metric and not any(word in metric for word in ('비율', '증감률')):
        return _format_people(number)
    return _compact(value, limit=18)


def _region_short_name(region_name: str) -> str:
    """긴 행정구역명은 본문 제목에서 마지막 시군구 이름으로 줄입니다."""
    parts = str(region_name or '').split()
    return parts[-1] if parts else '선택 지역'


def _cover_title_parts(title: str, timeframe: str) -> tuple[str, str, str]:
    """표지의 두 제목 상자가 서로 겹치지 않도록 전략명을 의미 단위로 나눕니다."""
    words = str(title or '지역 관광 실행전략').split()
    first_words: list[str] = []
    remainder: list[str] = []
    for word in words:
        candidate = ' '.join([*first_words, word])
        if not remainder and len(candidate) <= 18:
            first_words.append(word)
        else:
            remainder.append(word)
    first = ' '.join(first_words) or _compact(title, limit=18)
    remainder_text = ' '.join(remainder)
    remainder_text = re.sub(r'\s*(?:시범사업|사업|기획안)$', '', remainder_text).strip()
    second = (
        f'{_compact(remainder_text, limit=10)} 실행 기획안'
        if remainder_text else '실행 기획안'
    )
    return first, f'{timeframe} ', second


def _first_strategy(report: dict[str, Any]) -> dict[str, Any]:
    """현재 단일 기획안 구조에서 첫 전략을 안전하게 읽습니다."""
    strategies = list(report.get('strategies') or [])
    return strategies[0] if strategies else {}


def _finding(report: dict[str, Any], index: int, fallback: str) -> dict[str, str]:
    """관측지표 한 건을 카드·문장에 쓸 공통 형식으로 정리합니다."""
    findings = list(report.get('observed_findings') or [])
    item = findings[index] if index < len(findings) else {}
    return {
        'metric': _compact(item.get('metric') or fallback, limit=34),
        'value': _compact(item.get('value') or '자료 없음', limit=28),
        'interpretation': _compact(item.get('interpretation') or '공식 원자료 관측값', limit=90),
        'source': _compact(item.get('source') or '한국관광 데이터랩', limit=80),
    }


def _find_shape(slide: Any, name: str) -> Any:
    """템플릿의 안정적인 개체 이름으로 편집 대상을 찾습니다."""
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise ValueError(f'12장 PPT 템플릿 개체를 찾지 못했습니다: {name}')


def _copy_run_font(source: Any, target: Any) -> None:
    """새 문장에도 템플릿의 글꼴·크기·색상·굵기를 그대로 복사합니다."""
    target.name = source.name
    target.size = source.size
    target.bold = source.bold
    target.italic = source.italic
    target.underline = source.underline
    color_type = getattr(source.color, 'type', None)
    if color_type is None:
        return
    if getattr(source.color, 'rgb', None) is not None:
        target.color.rgb = source.color.rgb
    elif getattr(source.color, 'theme_color', None) is not None:
        target.color.theme_color = source.color.theme_color


def _copy_paragraph_style(source: Any, target: Any) -> None:
    """문단을 추가할 때 기존 정렬·간격·들여쓰기를 유지합니다."""
    target.alignment = source.alignment
    target.level = source.level
    target.line_spacing = source.line_spacing
    target.space_before = source.space_before
    target.space_after = source.space_after
    if source._p.pPr is not None:
        if target._p.pPr is not None:
            target._p.remove(target._p.pPr)
        target._p.insert(0, deepcopy(source._p.pPr))


def _set_lines(shape: Any, lines: Iterable[str]) -> None:
    """기존 텍스트 개체의 서식을 유지하면서 문단별 내용만 교체합니다."""
    if not getattr(shape, 'has_text_frame', False):
        raise ValueError(f'텍스트 개체가 아닙니다: {shape.name}')
    values = [str(value or '') for value in lines]
    if not values:
        values = ['']
    frame = shape.text_frame
    source_paragraphs = list(frame.paragraphs)
    source_paragraph = next((item for item in source_paragraphs if item.runs), source_paragraphs[0])
    source_font = source_paragraph.runs[0].font if source_paragraph.runs else None
    while len(frame.paragraphs) < len(values):
        paragraph = frame.add_paragraph()
        _copy_paragraph_style(source_paragraph, paragraph)
        run = paragraph.add_run()
        if source_font is not None:
            _copy_run_font(source_font, run.font)
    for index, paragraph in enumerate(frame.paragraphs):
        value = values[index] if index < len(values) else ''
        if paragraph.runs:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]:
                run.text = ''
        else:
            run = paragraph.add_run()
            run.text = value
            if source_font is not None:
                _copy_run_font(source_font, run.font)


def _set_text(slide: Any, name: str, value: Any) -> None:
    """이름으로 찾은 텍스트 개체의 문장만 바꿉니다."""
    _set_lines(_find_shape(slide, name), str(value or '').split('\n'))


def _set_two_tone(slide: Any, name: str, first: str, second: str) -> None:
    """제목의 네이비·시안 두 색 구성을 유지하며 두 구간만 교체합니다."""
    shape = _find_shape(slide, name)
    paragraph = shape.text_frame.paragraphs[0]
    while len(paragraph.runs) < 2:
        source_font = paragraph.runs[0].font if paragraph.runs else None
        run = paragraph.add_run()
        if source_font is not None:
            _copy_run_font(source_font, run.font)
    paragraph.runs[0].text = first
    paragraph.runs[1].text = second
    for run in paragraph.runs[2:]:
        run.text = ''
    for extra in shape.text_frame.paragraphs[1:]:
        for run in extra.runs:
            run.text = ''


def _source_note(source: dict[str, Any]) -> str:
    """사진·주장 출처를 발표자 노트의 한 줄 형식으로 바꿉니다."""
    title = _compact(source.get('title') or source.get('source_id') or '공식 자료', limit=90)
    url = str(source.get('source_url') or '').strip()
    return f'- {title}: {url}' if url else f'- {title}'


def _set_notes(slide: Any, sources: Iterable[dict[str, Any]], method: str = '') -> None:
    """외부 사진과 비자명한 주장 출처를 슬라이드 노트에 남깁니다."""
    lines = ['[Sources]']
    seen: set[tuple[str, str]] = set()
    for source in sources:
        key = (str(source.get('title') or source.get('source_id') or ''), str(source.get('source_url') or ''))
        if key in seen:
            continue
        seen.add(key)
        lines.append(_source_note(source))
    if len(lines) == 1:
        lines.append('- 보고서에 포함된 검증된 원자료·모델 메타데이터')
    if method:
        lines.extend(['', '[Method]', method])
    slide.notes_slide.notes_text_frame.text = '\n'.join(lines)


def _normalize_png(stream: BytesIO) -> tuple[bytes, float]:
    """외부 사진을 PPT 내부 이미지 형식인 PNG로 정규화합니다."""
    stream.seek(0)
    with Image.open(stream) as raw:
        image = ImageOps.exif_transpose(raw)
        if image.mode not in {'RGB', 'RGBA'}:
            image = image.convert('RGB')
        ratio = image.width / max(image.height, 1)
        output = BytesIO()
        image.save(output, format='PNG', optimize=True)
    return output.getvalue(), ratio


def _replace_picture(slide: Any, name: str, stream: BytesIO) -> None:
    """사진의 위치·마스크를 유지하고 내부 이미지 데이터와 자르기만 바꿉니다."""
    picture = _find_shape(slide, name)
    try:
        relationship_id = picture._element.blipFill.blip.rEmbed
        image_part = slide.part.related_part(relationship_id)
    except (AttributeError, KeyError) as exc:
        raise ValueError(f'사진 개체를 교체할 수 없습니다: {name}') from exc
    blob, image_ratio = _normalize_png(stream)
    image_part._blob = blob
    frame_ratio = picture.width / max(picture.height, 1)
    picture.crop_left = picture.crop_right = picture.crop_top = picture.crop_bottom = 0
    if image_ratio > frame_ratio:
        crop = max(0.0, (1 - frame_ratio / image_ratio) / 2)
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        crop = max(0.0, (1 - image_ratio / frame_ratio) / 2)
        picture.crop_top = crop
        picture.crop_bottom = crop


def _replace_template_photos(prs: Presentation, report: dict[str, Any]) -> dict[int, list[dict[str, Any]]]:
    """공식 Open API 사진을 최대 6장 사용하고 부족하면 같은 지역 사진을 순환합니다."""
    downloaded = download_images(select_image_sources(report, limit=6))
    targets = {
        3: 'tourism-photo',
        4: 'city-photo',
        6: 'visitor-photo',
        9: 'palace-photo',
        10: 'island-photo',
        11: 'workspace-photo',
    }
    notes_by_slide: dict[int, list[dict[str, Any]]] = {number: [] for number in targets}
    if not downloaded:
        # 공식 지역 사진이 없으면 강남 샘플 사진을 남기지 않고 중립적인 작업공간 이미지를 재사용합니다.
        fallback_picture = _find_shape(prs.slides[10], 'workspace-photo')
        fallback_part = prs.slides[10].part.related_part(fallback_picture._element.blipFill.blip.rEmbed)
        fallback_stream = BytesIO(fallback_part.blob)
        downloaded = [(
            {
                'title': '공식 지역 사진 미연결 시 사용하는 중립 템플릿 이미지',
                'source_type': 'template',
                'source_url': '',
            },
            fallback_stream,
        )]
    for index, (slide_number, shape_name) in enumerate(targets.items()):
        source, stream = downloaded[index % len(downloaded)]
        _replace_picture(prs.slides[slide_number - 1], shape_name, stream)
        notes_by_slide[slide_number].append(source)
    return notes_by_slide


def _scenario_rows(report: dict[str, Any]) -> dict[str, Any] | None:
    """저장 ML 자연추세와 사용자 입력 목표 수준을 같은 월 기준으로 계산합니다."""
    ml = report.get('ml_analysis') or {}
    selection = select_report_forecast(report)
    forecasts = selection['rows']
    if not forecasts:
        return None
    horizon = len(forecasts)
    target = execution_target(report)
    # 부분 전망의 마지막 달을 전체 사업의 최종 목표월로 바꾸지 않습니다.
    has_target = target is not None and selection['complete']
    visitor_target_pct, spending_target_pct = target if has_target else (0.0, 0.0)
    baseline_visitors = [_as_float(item.get('visitors')) for item in forecasts]
    baseline_spending = [_as_float(item.get('spending_krw')) for item in forecasts]
    target_visitors = target_series(baseline_visitors, visitor_target_pct)
    target_spending = target_series(baseline_spending, spending_target_pct)
    linked = visitor_linked_spending(baseline_visitors, baseline_spending, target_visitors)
    # 별도로 입력한 소비 목표는 보존합니다. 같은 목표율은 방문 증가분으로 산출합니다.
    linked_mode = has_target and visitor_target_pct == spending_target_pct and all(
        value is not None for value in linked['targets'])
    if linked_mode:
        target_spending = linked['targets']
    return {
        'categories': [_month_label(item.get('month')) for item in forecasts],
        'baseline_visitors': baseline_visitors,
        'baseline_spending': baseline_spending,
        'target_visitors': target_visitors,
        'target_spending': target_spending,
        'spending_per_visit_proxy': linked['ratios'],
        'spending_calculation': 'visitor_linked' if linked_mode else 'independent_target',
        'visitor_target_pct': visitor_target_pct,
        'spending_target_pct': spending_target_pct,
        'visitor_gap': sum(target_visitors) - sum(baseline_visitors),
        'spending_gap': sum(target_spending) - sum(baseline_spending),
        'horizon': horizon,
        'has_target': has_target,
        'source_period': str(ml.get('source_period') or ''),
        'period': selection['period'],
        'notes': selection['notes'],
    }


def _update_scenario_chart(slide: Any, scenario: dict[str, Any] | None) -> None:
    """슬라이드 9의 네이티브 차트를 실제 자연추세·목표 데이터로 교체합니다."""
    chart_shape = _find_shape(slide, 'visitor-natural-target-chart')
    if not getattr(chart_shape, 'has_chart', False):
        raise ValueError('12장 PPT 템플릿의 ML 시나리오 차트가 편집 가능한 차트가 아닙니다.')
    chart_data = ChartData()
    if scenario:
        chart_data.categories = scenario['categories']
        chart_data.add_series('ML 자연추세', [value / 1_000_000 for value in scenario['baseline_visitors']])
        if scenario['has_target']:
            chart_data.add_series('검토 목표', [value / 1_000_000 for value in scenario['target_visitors']])
    else:
        chart_data.categories = ['지원 전']
        chart_data.add_series('ML 자연추세', [0])
    chart = chart_shape.chart
    chart.replace_data(chart_data)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    colors = (NAVY, ORANGE)
    for index, series in enumerate(chart.series):
        series.format.line.color.rgb = colors[index % len(colors)]
        series.format.line.width = 19050
        series.format.fill.background()


def _budget_condition(report: dict[str, Any]) -> str:
    """담당자가 입력한 희망 예산 조건을 짧게 표시합니다."""
    brief = report.get('planning_brief') or {}
    status = str(brief.get('budget_status') or 'unknown')
    minimum = int(_as_float(brief.get('budget_min_krw')))
    maximum = int(_as_float(brief.get('budget_max_krw')))

    def amount_label(value: int) -> str:
        if value >= 100_000_000:
            return f'{value / 100_000_000:g}억 원'
        if value >= 10_000:
            return f'{value / 10_000:g}만 원'
        return f'{value:,}원'

    if status == 'unknown' or not maximum:
        return '희망 예산 미정'
    amount = (
        f'{amount_label(minimum)}~{amount_label(maximum)}'
        if minimum and minimum != maximum
        else amount_label(maximum)
    )
    suffix = ' · 상한' if brief.get('budget_hard_limit') else ''
    return _compact(f'희망 {amount}{suffix}', limit=45)


def _business_components(strategy: dict[str, Any]) -> list[tuple[str, str]]:
    """5단계 실행안·예산·KPI를 7개 사업 구성으로 재배열합니다."""
    steps = list(strategy.get('implementation_steps') or [])[:5]
    components = [
        (_compact(item.get('task') or f'{index + 1}단계 실행', limit=72),
         _compact(item.get('deliverable') or '확인 가능한 산출물', limit=58))
        for index, item in enumerate(steps)
    ]
    while len(components) < 5:
        components.append((f'{len(components) + 1}단계 실행', '확인 가능한 산출물'))
    components.extend([
        ('성과 측정', _compact(strategy.get('kpi') or '운영 전후 같은 기준의 KPI를 확인', limit=72)),
        ('예산 확정', _compact(strategy.get('budget') or '수량×단가 또는 비교견적으로 산정', limit=72)),
    ])
    return components[:7]


def _all_sources(report: dict[str, Any]) -> list[dict[str, Any]]:
    """발표자 노트에 기록할 검증된 출처 목록을 반환합니다."""
    return [item for item in (report.get('evidence_sources') or []) if isinstance(item, dict)]


def _source_domain(value: Any) -> str:
    """긴 URL을 출처 설명에 쓰기 쉬운 도메인으로 줄입니다."""
    try:
        return urlparse(str(value or '')).netloc or ''
    except ValueError:
        return ''


def _populate_cover(prs: Presentation, report: dict[str, Any]) -> None:
    """1~2장 표지와 목차를 현재 지역·전략명으로 채웁니다."""
    strategy = _first_strategy(report)
    region = str(report.get('region_name') or '선택 지역')
    timeframe = _compact(strategy.get('timeframe') or '3~6개월', limit=14)
    title_line, title_accent, title_tail = _cover_title_parts(
        str(strategy.get('title') or f'{region} 관광 전략'),
        timeframe,
    )
    offline = report.get('generation_mode') == 'offline_sample'
    slide = prs.slides[0]
    _set_text(slide, 'security-label', review_label(report))
    _set_text(slide, 'cover-title-line1', title_line)
    _set_two_tone(slide, 'cover-title-line2', title_accent, title_tail)
    _set_text(slide, 'cover-date', date.today().strftime('%Y년 %m월 %d일'))
    _set_text(slide, 'org-title', region)
    _set_text(slide, 'org-owner', '화면·출력 검토용 / Agent 미실행' if offline else '관광정책 검토용 / 데이터·AI 기반')
    _set_text(slide, 'report-number', f'보 고 서 출 력 일  {date.today():%Y-%m-%d}')
    _set_text(slide, 'report-topic', _compact(report.get('summary'), limit=28))

    toc = prs.slides[1]
    _set_text(toc, 'top-brand', f'{region} 관광 전략기획 데이터 시스템')
    _set_text(toc, 'top-date', date.today().strftime('%Y. %m. %d'))
    labels = [
        '데이터 근거와 기획 배경', '문제 정의 및 시범사업 목표', '참여자와 이용 시나리오',
        '핵심 사업 구성 7가지', '실행 Flow·산출물', '유사 지역 비교와 차별점',
        'ML 자연추세·목표 시나리오', '예산 산식·성과 측정',
        '5-Agent·공식 근거 파이프라인', '기대효과·확대 판단',
    ]
    for column in range(2):
        for row in range(5):
            _set_text(toc, f'toc-{column}-{row}-label', labels[column * 5 + row])


def _populate_evidence(prs: Presentation, report: dict[str, Any]) -> None:
    """3~4장에 관측 근거, 문제 정의, 시범 목표를 배치합니다."""
    strategy = _first_strategy(report)
    region = str(report.get('region_name') or '선택 지역')
    first = _finding(report, 0, '월간 방문자 수')
    second = _finding(report, 1, '관광소비액')
    third = _finding(report, 2, '숙박·체류 지표')

    slide = prs.slides[2]
    _set_two_tone(slide, 'title', f'{_region_short_name(region)} 관측 근거, ', '전환 과제를 가리키다')
    intro = _sentences(strategy.get('comparison_analysis'), limit=2, length=125)
    if not intro:
        intro = [
            f"{first['metric']} {first['value']}, {second['metric']} {second['value']}가 관측되었습니다.",
            _compact(strategy.get('problem_to_solve') or third['interpretation'], limit=125),
        ]
    _set_lines(_find_shape(slide, 'intro'), intro[:2])
    _set_text(slide, 'decision-title', f"{first['metric']} · {second['metric']}")
    _set_text(slide, 'decision-body', f"{first['value']} · {second['value']}\n{first['interpretation']}")
    _set_text(slide, 'activation-title', f"전환 점검: {third['metric']}")
    _set_text(slide, 'activation-body', f"{third['value']}\n{_compact(strategy.get('problem_to_solve'), limit=82)}")
    _set_text(slide, 'release-label', first['metric'])
    _set_text(slide, 'release-value', _metric_card_value(first['metric'], first['value']))
    _set_text(slide, 'release-suffix', 'LATEST')
    _set_text(slide, 'strategy-label', third['metric'])
    _set_text(slide, 'strategy-value', _metric_card_value(third['metric'], third['value']))
    _set_text(slide, 'strategy-suffix', 'CHECK')

    problem = prs.slides[3]
    _set_lines(_find_shape(problem, 'title'), ['문제 정의 및', '시범사업 목표'])
    features = [
        first['metric'], second['metric'], third['metric'],
        _compact(strategy.get('timeframe') or '3~6개월 검증', limit=14),
        '대상 설정', '현장 운영', '성과 측정', '확대 판단',
    ]
    for index, value in enumerate(features):
        _set_text(problem, f'feature-{index}', _compact(value, limit=16))
    body_lines = [
        _compact(strategy.get('problem_to_solve') or '관측 자료로 핵심 문제를 확인합니다.', limit=180),
        '',
        _compact(
            f"{strategy.get('solution') or '실행 가능한 시범사업을 설계합니다.'} "
            f"{strategy.get('expected_effect') or '운영 결과로 중단·보완·확대를 판단합니다.'}",
            limit=300,
        ),
    ]
    _set_lines(_find_shape(problem, 'body'), body_lines)


def _populate_people_and_components(prs: Presentation, report: dict[str, Any]) -> None:
    """5~6장에 참여자 역할과 7개 핵심 사업 구성을 채웁니다."""
    strategy = _first_strategy(report)
    components = _business_components(strategy)
    slide = prs.slides[4]
    _set_text(slide, 'subtitle', _compact(strategy.get('solution'), limit=110))
    audience = [
        ('목표 방문자', '기획안이 정한 동선·혜택을 이용하고 이용·결제·체류 반응을 남깁니다.', '행동: 방문 → 이용 → 선택 결과 확인'),
        ('지자체·운영기관', '참여 기준, 운영 품질, 예산 집행, 월별 KPI와 중단·보완·확대 판단을 관리합니다.', '역할: 기준 관리·성과 측정·의사결정'),
        ('관광·상권 사업자', '현장 상품과 혜택을 운영하고 정해진 항목으로 이용·결제·민원 데이터를 협력합니다.', '역할: 현장 운영·데이터 협력'),
    ]
    for index, (title, body, role) in enumerate(audience):
        _set_text(slide, f'audience-pill-{index}', title)
        _set_text(slide, f'audience-body-{index}', body)
        _set_text(slide, f'audience-role-{index}', role)
    _set_text(slide, 'footer-left', f"{strategy.get('title') or '지역 관광 전략'} | {strategy.get('timeframe') or '3~6개월'}")
    _set_text(slide, 'footer-right', 'Evidence-Driven Tourism Plan')

    functions = prs.slides[5]
    for index in range(3):
        title, body = components[index]
        _set_text(functions, f'function-label-{index}', f'{title}: {body}')
    left = '  '.join(f'④ {components[3][0]}: {components[3][1]}  ⑤ {components[4][0]}: {components[4][1]}'.split())
    right = '  '.join(f'⑥ {components[5][0]}: {components[5][1]}  ⑦ {components[6][0]}: {components[6][1]}'.split())
    _set_text(functions, 'function-more-left', left)
    _set_text(functions, 'function-more-right', right)


def _populate_timeline(prs: Presentation, report: dict[str, Any]) -> None:
    """7장에 정확히 5개 실행 단계와 산출물을 표시합니다."""
    strategy = _first_strategy(report)
    steps = list(strategy.get('implementation_steps') or [])[:5]
    while len(steps) < 5:
        steps.append({'step': len(steps) + 1, 'schedule': '', 'task': '실행 단계', 'deliverable': '확인 산출물'})
    slide = prs.slides[6]
    _set_two_tone(slide, 'title', f"{strategy.get('timeframe') or '3~6개월'} 실행 ", 'Flow 및 산출물')
    for index, item in enumerate(steps[:4]):
        _set_text(slide, f'step-{index}-label', f"STEP {index + 1:02d} · {_compact(item.get('schedule'), limit=12)}")
        _set_text(slide, f'step-{index}-title', _compact(item.get('task'), limit=32))
        _set_text(
            slide,
            f'step-{index}-body',
            f"{_compact(item.get('task'), limit=94)}\n산출물: {_compact(item.get('deliverable'), limit=58)}",
        )
    last = steps[4]
    _set_text(slide, 'value-title', f"STEP 05 · {_compact(last.get('schedule') or '최종 판단', limit=25)}")
    _set_text(
        slide,
        'value-body',
        f"{_compact(last.get('task'), limit=180)}\n산출물: {_compact(last.get('deliverable'), limit=95)}. "
        '목표 달성 여부뿐 아니라 비용·이탈·현장 운영 가능성을 함께 보고 중단·보완·확대를 결정합니다.',
    )


def _populate_comparison(prs: Presentation, report: dict[str, Any]) -> None:
    """8장에 관측·비교·기획 판단을 네 개 근거 블록으로 정리합니다."""
    strategy = _first_strategy(report)
    findings = [_finding(report, index, f'관측지표 {index + 1}') for index in range(3)]
    comparison_sentences = _sentences(strategy.get('comparison_analysis'), limit=2, length=150)
    blocks = [
        (findings[0]['metric'], f"{findings[0]['value']}. {findings[0]['interpretation']}", '관측 사실에서 출발'),
        (findings[1]['metric'], f"{findings[1]['value']}. {findings[1]['interpretation']}", '같은 기준으로 확인'),
        ('비교 지역 신호', comparison_sentences[0] if comparison_sentences else '검증된 비교 자료가 부족해 선택 지역의 월별 변화만 사용합니다.', '비교는 인과효과가 아님'),
        ('실행 판단', _compact(strategy.get('problem_to_solve'), limit=150), _compact(strategy.get('timeframe') or '작게 검증 후 확대', limit=30)),
    ]
    slide = prs.slides[7]
    for index, (title, body, plus) in enumerate(blocks):
        _set_text(slide, f'diff-title-{index}', title)
        _set_text(slide, f'diff-body-{index}', body)
        _set_text(slide, f'diff-plus-{index}', f'+ {plus}')


def _populate_ml(prs: Presentation, report: dict[str, Any]) -> None:
    """9장에 자연추세, 사용자 목표, 차이의 정확한 의미를 표시합니다."""
    slide = prs.slides[8]
    scenario = _scenario_rows(report)
    _update_scenario_chart(slide, scenario)
    if not scenario:
        _set_two_tone(slide, 'title', '저장 ML 전망 ', '미지원')
        _set_lines(_find_shape(slide, 'intro'), [
            '이 지역은 검증된 저장 ML 전망이 없어 수치 목표 비교를 만들지 않습니다.',
            '관측 원자료와 공식 근거를 사용한 실행 기획은 다른 장에서 계속 확인할 수 있습니다.',
        ])
        _set_text(slide, 'prediction-title', '자연추세 전망 상태')
        _set_text(slide, 'prediction-body', ' '.join(select_report_forecast(report)['notes']))
        _set_text(slide, 'rag-title', '기획안 적용 원칙')
        _set_text(slide, 'rag-body', '관측 사실과 공식 사례만 사용하고 성과 수치를 보장하지 않습니다.')
        _set_text(slide, 'forecast-chart-note', 'ML 미지원 · 정책효과 예측 아님')
        return
    horizon = scenario['horizon']
    _set_two_tone(slide, 'title', f'{horizon}개월 자연추세와 ', '검토 목표')
    target_label = (
        f"방문 +{scenario['visitor_target_pct']:.1f}% · 소비 +{scenario['spending_target_pct']:.1f}%"
        if scenario['has_target'] else '사용자 목표 미입력'
    )
    _set_lines(_find_shape(slide, 'intro'), [
        'ML은 기존 이력으로 자연추세를 계산합니다. 정책을 하지 않았을 때의 인과적 반사실이나 사업효과 예측이 아닙니다.',
        f'{target_label}은 담당자가 검토를 위해 입력한 목표 수준이며 자연추세와의 차이를 성과 보장으로 해석하지 않습니다.',
    ])
    natural_people = sum(scenario['baseline_visitors'])
    natural_spending = sum(scenario['baseline_spending'])
    _set_text(slide, 'prediction-title', f'{horizon}개월 ML 자연추세')
    _set_text(
        slide,
        'prediction-body',
        f"방문자 합계 {_format_people(natural_people)} · 관광소비 {_format_money(natural_spending)}. "
        f"기획 구간 {scenario['period']}의 월별 합계입니다(기간 순방문자 아님).",
    )
    _set_text(slide, 'rag-title', '자연추세와 목표 수준의 산술 차이')
    if scenario['has_target']:
        gap_text = f"방문 {_format_people(scenario['visitor_gap'])} · 소비 {_format_money(scenario['spending_gap'])}"
    else:
        gap_text = '담당자 목표율을 입력하면 필요한 월별 수준을 함께 계산합니다.'
    _set_text(slide, 'rag-body', f'{gap_text}. 이 차이는 정책 인과효과나 보장 성과가 아닙니다. ' + ' '.join(scenario['notes']))
    _set_text(slide, 'tech-tag-0', '시간순 검증')
    _set_text(slide, 'tech-tag-1', '저장 모델')
    _set_text(slide, 'tech-tag-2', f'{horizon}개월 전망')
    _set_text(
        slide,
        'forecast-chart-note',
        '단위: 백만 명 · 자연추세와 사용자 검토 목표'
        if scenario['has_target'] else '단위: 백만 명 · 사용자 목표 미입력',
    )


def _populate_budget(prs: Presentation, report: dict[str, Any]) -> None:
    """10장에 예산 산식과 운영 KPI를 근거 없는 총액 없이 표시합니다."""
    strategy = _first_strategy(report)
    slide = prs.slides[9]
    budget_sentences = _sentences(strategy.get('budget'), limit=2, length=140)
    if not budget_sentences:
        budget_sentences = ['수량×단가 또는 비교견적으로 예산을 산정합니다.', '희망 예산이 미정이면 범위와 우선순위부터 정합니다.']
    _set_lines(_find_shape(slide, 'intro'), budget_sentences[:2])
    _set_text(slide, 'dashboard-title', f"예산 조건 · {_budget_condition(report)}")
    _set_text(slide, 'dashboard-body', _compact(strategy.get('budget'), limit=135))
    _set_text(slide, 'analysis-title', '성과 판단 KPI')
    _set_text(slide, 'analysis-body', _compact(strategy.get('kpi'), limit=135))
    _set_text(slide, 'report-card-title', '예산 확정 원칙')
    _set_text(slide, 'report-card-body', '수량·단가·운영범위 검증 후')
    _set_text(slide, 'report-card-orange', '공식 단가·비교견적으로 확정')
    _set_text(slide, 'map-card-title', '측정 설계 원칙')
    _set_text(slide, 'map-card-body', '운영 전·후 같은 기준과')
    _set_text(slide, 'map-card-orange', '참여·비참여 집단을 구분')


def _populate_agents(prs: Presentation, report: dict[str, Any]) -> None:
    """11장에 실제 5-Agent와 MySQL·ML·RAG·OpenAI 역할을 표시합니다."""
    slide = prs.slides[10]
    offline = report.get('generation_mode') == 'offline_sample'
    trace = {
        str(item.get('agent') or item.get('name') or '').lower(): item
        for item in (report.get('agent_trace') or [])
        if isinstance(item, dict)
    }
    cards = [
        ('Evidence Agent', 'MySQL 실제값·ML 자연추세·지역 근거 정리'),
        ('Case Scout', 'RAG와 공식 웹에서 유사 사례·예산·성과 검색'),
        ('Transferability', '지역 차이와 적용·제외 조건, 시범 범위 검토'),
        ('Planner Agent', '문제·솔루션·5단계·예산 산식·KPI 설계'),
    ]
    for index, (title, body) in enumerate(cards):
        _set_text(slide, f'stack-card-{index}-title', title)
        if offline:
            status = '이번 샘플 미실행'
        elif trace:
            key = ('evidence', 'case_scout', 'transferability', 'planner')[index]
            status = '실행 기록 있음' if any(key in name for name in trace) else '실행 기록 미확인'
        else:
            status = '역할 정의'
        _set_text(slide, f'stack-card-{index}-body', f'{body}\n상태: {status}')
    sources = _all_sources(report)
    domains = [domain for domain in (_source_domain(item.get('source_url')) for item in sources) if domain]
    domain_text = ', '.join(dict.fromkeys(domains[:4])) or '검증된 원자료·모델 메타데이터'
    if offline:
        body_lines = [
            '⑤ Reviewer Agent는 실제 생성 시 관측 사실·ML 전망·사용자 목표·추천을 다시 분리하고 근거 없는 수치와 관광지를 차단합니다.',
            '',
            '이번 파일은 API 비용 없이 레이아웃과 데이터 배치를 검토하는 오프라인 샘플입니다. OpenAI와 5-Agent는 실행하지 않았습니다.',
            '',
            f'실제 생성에서는 MySQL 사실, ML 자연추세, RAG 공식 사례를 검증한 뒤 JSON 품질 게이트를 통과한 결과만 출력합니다. 연결 자료: {domain_text}',
        ]
    else:
        body_lines = [
            '⑤ Reviewer Agent가 관측 사실·ML 전망·사용자 목표·추천을 다시 분리하고 근거 없는 수치와 관광지를 차단합니다.',
            '',
            'OpenAI는 검증된 근거 묶음을 받은 뒤 문장 구성, 대안 비교, 실행안 초안을 담당합니다. MySQL은 사실, ML은 자연추세, RAG는 공식 사례, LLM은 설명과 기획을 맡습니다.',
            '',
            f'최종 JSON 스키마와 품질 게이트를 통과한 결과만 Word·PPT로 변환합니다. 주요 출처: {domain_text}',
        ]
    _set_lines(_find_shape(slide, 'body'), body_lines)


def _populate_decision(prs: Presentation, report: dict[str, Any]) -> None:
    """12장에 기대 변화와 중단·보완·확대 기준을 정리합니다."""
    strategy = _first_strategy(report)
    slide = prs.slides[11]
    _set_text(slide, 'expected-body', _compact(strategy.get('expected_effect'), limit=155))
    _set_text(slide, 'indicator-pill', '월별 KPI 추적')
    _set_text(slide, 'service-pill', f"{strategy.get('timeframe') or '3~6개월'} 판단 게이트")
    _set_text(
        slide,
        'service-body',
        '목표 달성 여부뿐 아니라 비용, 이탈 구간, 민원, 사업자 운영 가능성을 함께 보고 중단·보완·확대를 결정합니다.',
    )
    _set_text(slide, 'service-note', '* 자연추세와 실제 운영 결과를 비교하되 관측 설계가 약하면 인과효과로 표현하지 않습니다.')
    _set_text(slide, 'service-scope', '의사결정: 증거 우선')
    _set_text(slide, 'private-pill', '근거 확보 후 확장')
    _set_text(
        slide,
        'private-body',
        '검증된 동선·업종·계절부터 확대하고, 이후 외국인·MICE·민간 결제 등 필요한 데이터를 단계적으로 연계합니다.',
    )
    _set_text(slide, 'private-goal', '핵심: 성과 검증 후 예산 확대')
    _set_text(slide, 'private-scope', '원칙: 효과 보장 표현 금지')


def _apply_source_notes(
    prs: Presentation,
    report: dict[str, Any],
    photo_notes: dict[int, list[dict[str, Any]]],
) -> None:
    """각 장의 주장·사진 출처와 계산 방법을 발표자 노트에 기록합니다."""
    sources = _all_sources(report)
    strategy = _first_strategy(report)
    evidence_ids = {
        item.strip()
        for item in re.split(r'[,\s]+', str(strategy.get('evidence') or ''))
        if item.strip()
    }
    selected = [item for item in sources if str(item.get('source_id') or '') in evidence_ids]
    base_sources = selected or sources[:6]
    methods = {
        1: '표지의 지역·전략명·기간은 저장된 ReportResponse에서 읽음.',
        2: '목차는 승인된 12장 템플릿 구조를 고정 사용.',
        3: '관측값과 비교 문장은 보고서의 observed_findings·comparison_analysis를 사용.',
        4: '문제·솔루션·기대 변화는 단일 검증 기획안을 요약.',
        5: '참여자 역할은 기획안 실행 주체와 측정 역할을 구분해 정리.',
        6: '5단계 실행안과 예산·KPI를 7개 사업 구성으로 재배열.',
        7: 'implementation_steps의 일정·행동·산출물을 그대로 사용.',
        8: '지역 비교는 인과효과가 아닌 우선순위 판단 신호로만 사용.',
        9: '저장 ML 자연추세와 사용자 입력 목표율을 분리. 목표 차이는 산술 차이이며 정책효과 예측 아님.',
        10: '예산은 총액을 임의 생성하지 않고 수량×단가 또는 비교견적 산식으로 제시.',
        11: 'Evidence + Case Scout 병렬 → Transferability → Planner → Reviewer 고정 흐름.',
        12: '기대효과는 제안이며 실제 운영 후 같은 기준의 관측값으로 검증.',
    }
    for number, slide in enumerate(prs.slides, start=1):
        slide_sources = [*base_sources, *photo_notes.get(number, [])]
        _set_notes(slide, slide_sources, methods[number])


def _validate_presentation(prs: Presentation, report: dict[str, Any]) -> None:
    """다운로드 전 장수·필수 섹션·편집 가능한 차트·샘플 잔여 문구를 검사합니다."""
    if len(prs.slides) != 12:
        raise ValueError(f'PowerPoint는 표지 포함 정확히 12장이어야 합니다: {len(prs.slides)}장')
    all_text = '\n'.join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if hasattr(shape, 'text')
    )
    required = (
        '실행 기획안', '문제 정의', '핵심 사업 구성 7가지', '실행 Flow',
        '자연추세', '예산', 'Reviewer Agent', '확대 판단',
    )
    missing = [label for label in required if label not in all_text]
    if missing:
        raise ValueError(f"PowerPoint 필수 섹션 누락: {', '.join(missing)}")
    chart_shape = _find_shape(prs.slides[8], 'visitor-natural-target-chart')
    if not getattr(chart_shape, 'has_chart', False):
        raise ValueError('ML 시나리오 장의 차트가 편집 가능한 PowerPoint 차트가 아닙니다.')
    region = str(report.get('region_name') or '')
    if '강남' not in region and re.search(r'강남(?:구|\s|$)', all_text):
        raise ValueError('다른 지역 보고서에 강남구 샘플 문구가 남아 있습니다.')


def create_strategy_proposal_presentation(report: dict[str, Any]) -> BytesIO:
    """검증된 보고서 값으로 승인된 16:9, 12장 편집형 PPTX를 반환합니다."""
    if not PRESENTATION_TEMPLATE_PATH.is_file():
        raise ValueError(f'12장 PowerPoint 템플릿 파일이 없습니다: {PRESENTATION_TEMPLATE_PATH.name}')
    prs = Presentation(str(PRESENTATION_TEMPLATE_PATH))
    _populate_cover(prs, report)
    _populate_evidence(prs, report)
    _populate_people_and_components(prs, report)
    _populate_timeline(prs, report)
    _populate_comparison(prs, report)
    _populate_ml(prs, report)
    _populate_budget(prs, report)
    _populate_agents(prs, report)
    _populate_decision(prs, report)
    photo_notes = _replace_template_photos(prs, report)
    _apply_source_notes(prs, report, photo_notes)
    _validate_presentation(prs, report)
    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output
