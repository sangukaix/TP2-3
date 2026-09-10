"""Execution overview plus one plain-language implementation page."""
from .proposal_layout_v8 import *
from . import proposal_layout_v8 as v8


LABELS=('준비와 협의','운영 절차 점검','시범 운영','실적 확인과 평가')
ROLES=('사업 총괄','운영·정산 담당','현장 운영 담당','사업·평가 담당')


def execution_groups(report):
    steps=[v for v in base._first_strategy(report).get('implementation_steps') or [] if isinstance(v,dict)]
    groups=[[v] for v in steps[:3]]+[steps[3:]]
    while len(groups)<4:groups.append([])
    return groups


def period(group):
    return ' / '.join(dict.fromkeys(copy_text(v.get('schedule')) for v in group if v.get('schedule'))) or '일정 협의'


def roadmap(prs,report):
    s=prs.slides[6];header(s,'4단계 실행 가이드')
    groups=execution_groups(report)
    # Connected milestones are editable PowerPoint shapes, not a pasted image.
    rect(s,'flow-track',260,260,1065,3,BORDER)
    rows=[['단계','기간','담당 역할(제안)','세부 실행 업무']]
    for i,group in enumerate(groups):
        x=106+i*355;cx=x+158
        ring=s.shapes.add_shape(MSO_SHAPE.OVAL,(cx-29)*EMU,233*EMU,58*EMU,58*EMU)
        ring.name=f'flow-node-{i}';ring.fill.solid();ring.fill.fore_color.rgb=RGBColor.from_string(BLUE if i%2==0 else CYAN);ring.line.fill.background()
        centered(text(s,f'flow-number-{i}',f'{i+1:02d}',cx-28,242,56,39,25,WHITE,True))
        centered(text(s,f'step-{i}-label',LABELS[i],x,306,316,43,25,BLUE,True))
        centered(text(s,f'step-{i}-period',period(group),x,354,316,47,20,SLATE))
        tasks='\n'.join(copy_text(v.get('task')) for v in group) or '담당자와 작업 범위를 정합니다.'
        rows.append([f'{i+1:02d}',period(group),ROLES[i],tasks])
    sh=table(s,'execution-table',rows,106,435,1388,371,[95,220,245,828],21)
    for cell in sh.table.rows[0].cells:
        for p in cell.text_frame.paragraphs:p.alignment=PP_ALIGN.CENTER
    for row in list(sh.table.rows)[1:]:
        for cell in list(row.cells)[:3]:
            for p in cell.text_frame.paragraphs:p.alignment=PP_ALIGN.CENTER
    text(s,'roadmap-note','단계별 실행 방법은 다음 장에서 확인할 수 있습니다. 담당 역할과 일정은 착수 전에 확정합니다.',106,835,1388,44,19,SLATE)


def execution_copy(report):
    """Supplement saved tasks with proposed procedures, never invented agreements."""
    strategy=base._first_strategy(report)
    from .proposal_presentation_v4 import _selected_candidate
    candidate=_selected_candidate(report)
    mechanism=' '.join(str(v or '') for v in (strategy.get('title'),strategy.get('solution'),candidate.get('mechanism')))
    if any(w in mechanism for w in ('환급','반값')):
        return [
            ('참여 가게와 지원 기준을 먼저 정합니다.',
             '참여를 희망하는 가게에 사용 가능한 상품권과 정산 방법을 확인합니다. 지원 대상, 인정할 지출, 신청 마감일과 1인당 한도를 정해 같은 안내문을 배포합니다.'),
            ('신청부터 지급까지 직접 시험합니다.',
             '시험 신청서를 넣고 영수증의 날짜·가게·금액을 확인합니다. 취소 결제와 중복 신청을 걸러내고, 서류가 부족할 때 누구에게 보완을 요청할지도 정합니다.'),
            ('작게 시작하고 확인된 신청부터 지급합니다.',
             '참여 가게를 한정해 먼저 운영합니다. 담당자가 지출 증빙을 확인한 뒤 정해진 기준에 따라 환급하고, 신청자에게 지급일과 상품권 사용 방법을 알려줍니다.'),
            ('지급 기록과 실제 사용 결과를 확인합니다.',
             '매주 신청·승인·지급 건수와 비용을 맞춰 봅니다. 상품권 사용 집계가 가능한지 운영사에 확인하고, 종료 후 취소를 뺀 실적을 정리합니다. 전체 관광소비 증가를 전부 사업 성과로 보지는 않습니다.'),
        ]
    return [
        ('함께 운영할 곳과 업무 범위를 정합니다.',
         '계획에 나온 참여처와 운영 가능 시간을 확인합니다. 누가 이용자를 안내하고 문의에 답할지 정한 뒤, 신청 대상과 이용 조건을 하나의 안내문으로 정리합니다.'),
        ('이용자 입장에서 전 과정을 시험합니다.',
         '안내를 보고 신청한 뒤 현장에서 이용을 마칠 때까지 직접 따라가 봅니다. 예약이 안 되거나 취소가 생겼을 때의 처리 방법을 정하고 담당자에게 공유합니다.'),
        ('제한된 범위에서 먼저 운영합니다.',
         '협의된 참여처와 시간대로 시작합니다. 현장 담당자가 안내와 안전 상태를 확인하고, 매일 신청 수와 실제 이용 수를 기록합니다. 반복되는 문의는 안내문에 반영합니다.'),
        ('이용 실적과 비용을 함께 정리합니다.',
         '매주 이용 기록과 지출 내역을 확인하고 불편이 생긴 원인을 정리합니다. 종료 후 같은 기준의 운영 전 자료와 비교하되, 계절 변화와 사업 효과를 구분해 보고합니다.'),
    ]


def execution_detail(prs,report):
    from pptx.opc.packuri import PackURI
    existing={str(p.partname) for p in prs.part.package.iter_parts()}
    number=1
    while f'/ppt/slides/slide{number}.xml' in existing:number+=1
    s=prs.slides.add_slide(prs.slides[5].slide_layout)
    # Deleted template slides leave numbering holes; avoid add_slide's len+1 collision.
    s.part._partname=PackURI(f'/ppt/slides/slide{number}.xml')
    header(s,'세부 실행내역')
    text(s,'execution-intro',short_title(report)+'의 단계별 실행 방법입니다. 적용 조건은 담당자와 협의해 확정합니다.',106,204,1388,54,23,SLATE)
    for i,(lead,body) in enumerate(execution_copy(report)):
        y=284+i*134
        rect(s,f'detail-rule-{i}',106,y+120,1388,1,BORDER)
        text(s,f'detail-number-{i}',f'{i+1:02d}',108,y+3,70,55,35,CYAN,True)
        text(s,f'detail-stage-{i}',LABELS[i],194,y+7,287,48,25,BLUE,True)
        text(s,f'detail-role-{i}',ROLES[i],194,y+55,287,39,19,SLATE)
        text(s,f'detail-lead-{i}',lead,514,y,966,41,24,INK,True)
        text(s,f'detail-body-{i}',body,514,y+45,966,73,22,SLATE)
    s.notes_slide.notes_text_frame.text='실행 방법 제안. 저장된 사업 내용에 맞춰 작성한 운영 절차이며 확정 협약이나 공식 지침이 아닙니다.\n'+str(base._first_strategy(report))
    sid=prs.slides._sldIdLst[-1];prs.slides._sldIdLst.remove(sid);prs.slides._sldIdLst.insert(6,sid)


def finish(prs,report):
    v8.finish(prs,report)
    execution_detail(prs,report)
    labels=['프로젝트 개요','사업 설계와 3개월 목표','지역별 참고 사례','사례 선정 근거','4단계 실행 가이드','세부 실행내역','ML 전망과 목표 KPI','사례 실적과 목표 설정','견적','근거·데이터·출처','맺음말']
    # Rebuild the text list because v8 removed the ninth menu item.
    toc=prs.slides[1]
    for sh in list(toc.shapes):
        if sh.name.startswith('toc-'):sh._element.getparent().remove(sh._element)
    for i,label in enumerate(labels):
        col,row=divmod(i,6);x=562+540*col;y=296+80*row
        text(toc,f'toc-{col}-{row}-number',f'{i+1:02d}',x,y,61,44,25,CYAN,True)
        text(toc,f'toc-{col}-{row}-label',label,x+65,y+5,415,55,25,INK,True)
        rect(toc,f'toc-{col}-{row}-rule',x,y+64,477,1,BORDER)


def target_evidence_page(prs, report):
    """The published outcome and our target assumption are shown separately."""
    from pptx.opc.packuri import PackURI
    basis = report['target_proposal_basis']
    existing = {str(p.partname) for p in prs.part.package.iter_parts()}
    number = 1
    while f'/ppt/slides/slide{number}.xml' in existing:
        number += 1
    s = prs.slides.add_slide(prs.slides[3].slide_layout)
    s.part._partname = PackURI(f'/ppt/slides/slide{number}.xml')
    header(s, '사례 실적과 목표 KPI 설정 근거')
    text(s, 'basis-case-title', basis['case_title'], 106, 205, 1388, 52, 30, BLUE, True)
    if basis['observed_visitors'] is not None:
        rows = [['비교 항목', '관광객 수', '값의 성격'],
                ['2023년 비교 규모', '약 198.4만 명', '248만 ÷ 1.25 역산'],
                ['2024년 10월 말 누적', '약 248만 명', '강진군의회 발표'],
                ['전년 대비 증가', '약 49.6만 명 / 25%', '발표 증가율 · 차이는 역산']]
        table(s, 'case-result-table', rows, 106, 279, 1388, 246, [455, 390, 543], 24)
    else:
        text(s, 'basis-description', '선정 사례의 운영 방식을 지역의 시범사업에 적용합니다. 비교 가능한 전후 실적이 없는 경우 목표는 계획 가정으로 제시합니다.', 106, 285, 1388, 170, 27, SLATE)
    scenario = report['execution_scenario']
    text(s, 'proposed-target', f"최종월 계획 목표  방문 +{scenario['visitor_target_pct']:g}% · 소비 +{scenario['spending_target_pct']:g}%", 106, 557, 1388, 63, 33, ORANGE, True)
    text(s, 'basis-method', basis['explanation'], 106, 637, 1388, 165, 21, SLATE)
    source_shape = text(s, 'basis-source', '출처: 강진군의회 제307회 본회의 시정연설, 2024.11.20' if basis['source_url'] else '목표율: 조정 가능한 기획 가정', 106, 825, 1388, 42, 19, SLATE)
    if basis['source_url']:
        for paragraph in source_shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.hyperlink.address = basis['source_url']
    s.notes_slide.notes_text_frame.text = basis['explanation'] + '\n' + basis['source_url']
    sid = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(sid)
    prs.slides._sldIdLst.insert(8, sid)


def case_selection_basis(report):
    """Explain the selected source linkage, without inventing environmental similarity."""
    from .proposal_presentation_v4 import _selected_candidate
    candidate=_selected_candidate(report)
    decision=report.get('planning_decision') or {}
    ids=candidate.get('case_source_ids') or decision.get('recommended_case_ids') or []
    sources=[source for source in report.get('evidence_sources') or []
             if source.get('source_id') in ids and source.get('source_type')=='benchmark_case']
    source=next(iter(sources),{})
    refund=any(word in str(source.get('operating_model') or '') for word in ('환급','돌려받')) and '상품권' in str(source.get('operating_model') or '')
    facts=[f"{row.get('metric')}: {row.get('value')}" for row in report.get('observed_findings') or []
           if row.get('metric') and row.get('value')][:2]
    return {'source':source, 'sources':sources, 'candidate':candidate,
            'reason':'소비 혜택이 다시 지역 상점 이용으로 이어지는 구조를 이번 사업의 핵심으로 삼아 참고했습니다.' if refund else '이번 기획안의 실행 방법을 구체화하기 위해, 연결된 공식 사례의 운영 절차를 참고했습니다.',
            'facts':' · '.join(facts) or '저장 보고서의 지역 현황을 바탕으로 운영 범위를 정합니다.',
            'operation':source.get('operating_model') or source.get('intervention') or '선정 사례의 운영 방식은 연결된 공식 원문을 기준으로 설명합니다.',
            'adaptation':candidate.get('mechanism') or base._first_strategy(report).get('solution') or '선택 지역의 사업 대상과 참여 범위에 맞춰 설계합니다.'}


def case_selection_page(prs, report):
    from pptx.opc.packuri import PackURI
    import json
    basis=case_selection_basis(report);source=basis['source']
    existing={str(part.partname) for part in prs.part.package.iter_parts()}
    number=1
    while f'/ppt/slides/slide{number}.xml' in existing:number+=1
    s=prs.slides.add_slide(prs.slides[4].slide_layout)
    s.part._partname=PackURI(f'/ppt/slides/slide{number}.xml')
    header(s,'참고 사례 선정 근거')
    case_name=' · '.join(filter(None,[source.get('case_region'),source.get('intervention') or source.get('title')]))
    text(s,'selection-case',case_name or '현재 기획안에 연결된 참고 사례',106,206,1388,85,32,BLUE,True)
    text(s,'selection-intro',basis['reason'],106,295,1388,70,25,SLATE)
    rows=[('기획 지역의 관측 현황',basis['facts']),
          ('공식 사례에서 참고한 방식',basis['operation']),
          ('이번 기획안에 적용한 부분',basis['adaptation'])]
    for i,(label,body) in enumerate(rows):
        y=387+i*113
        text(s,f'selection-label-{i}',label,106,y,330,88,24,BLUE,True)
        text(s,f'selection-body-{i}',copy_text(body),464,y,1030,99,24,INK)
        rect(s,f'selection-rule-{i}',106,y+103,1388,1,BORDER)
    text(s,'selection-limit','해석 범위 · 이 사례의 선택은 인구·자연환경·교통이 같은 지역이라는 뜻이 아닙니다. 환경 유사성과 동일한 성과율을 입증한 비교는 아닙니다.',106,745,1388,67,21,SLATE)
    link=text(s,'selection-source','출처 · '+str(source.get('title') or '저장 기획안의 사례 연결 정보'),106,824,1388,48,20,BLUE)
    if str(source.get('source_url') or '').startswith(('http://','https://')):
        for paragraph in link.text_frame.paragraphs:
            for run in paragraph.runs:run.hyperlink.address=source['source_url']
    s.notes_slide.notes_text_frame.text=json.dumps(basis,ensure_ascii=False,indent=2)
    sid=prs.slides._sldIdLst[-1];prs.slides._sldIdLst.remove(sid);prs.slides._sldIdLst.insert(5,sid)
