"""Evidence-first presentation finishing pass, with editable native content."""
import json
from pathlib import Path
from .proposal_layout_v9 import *
from .case_recommendation import report_cases, case_reference_role, VERSION

ASSETS = Path(__file__).resolve().parents[1] / 'assets' / 'proposal_cases'


def case_narrative(value):
    # Web citations retain their actual URL in the source link and speaker notes.
    # A raw Markdown URL is not prose and can force case cards into tiny type.
    return re.sub(r'\s*\(?\[[^\]]+\]\(https?://[^\s]+\)\)?', '', str(value or '')).strip()


def safe_text(slide,name,value,x,y,w,h,size=22,color=SLATE,bold=False):
    shape=text(slide,name,value,x,y,w-40,h,size,color,bold)
    shape.width=w*EMU
    return shape


def link(shape, url):
    if str(url or '').startswith(('https://', 'http://')):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs: run.hyperlink.address = url


def case_cards(slide, report):
    from .case_images import case_image
    from PIL import Image
    header(slide, '지역별 참고 사례')
    rows, primary = report_cases(report)
    image_sources=[]
    text(slide,'case-intro','운영 방식이 맞는 공식 사례를 연결합니다. 지역 자체의 유사도 순위와는 구분합니다.',106,188,1388,47,22,SLATE)
    for i, source in enumerate(rows):
        x=106+i*472
        rect(slide,f'case-panel-{i}',x,258,444,542,PALE if source in primary else WHITE)
        text(slide,f'case-number-{i}',f'{i+1:02d}',x+24,277,64,44,29,CYAN,True)
        text(slide,f'case-badge-{i}',case_reference_role(report, source),x+99,285,320,33,19,BLUE,True)
        name=(source.get('case_region') or '')+' · '+(source.get('intervention') or source.get('title') or '')
        safe_text(slide,f'case-title-{i}',name,x+24,339,396,98,25,INK,True)
        safe_text(slide,f'case-body-{i}',case_narrative(source.get('operating_model')),x+24,452,396,158,21,SLATE)
        image_info=case_image(source)
        if image_info:
            asset=image_info['path']
            with Image.open(asset) as raw:
                scale=min(396/raw.width,180/raw.height)
                w,h=raw.width*scale,raw.height*scale
            picture=slide.shapes.add_picture(str(asset),int((x+24+(396-w)/2)*EMU),int((615+(180-h)/2)*EMU),int(w*EMU),int(h*EMU))
            picture.name=f'case-photo-{i}'
            link_caption=text(slide,f'case-image-caption-{i}',image_info['caption']+'\n출처: '+image_info['credit'],x+24,805,396,53,16,SLATE)
            link(link_caption,image_info['page_url'])
            image_sources.append({k:str(v) for k,v in image_info.items() if k!='path'})
        else:
            text(slide,f'case-image-caption-{i}','해당 사례의 확인된 이미지 미등록\n아래 공식 사업 원문에서 확인할 수 있습니다.',x+24,705,396,80,18,SLATE)
        shape=text(slide,f'case-source-{i}','사업 원문 보기 ↗',x+24,868,396,24,16,BLUE)
        link(shape,source.get('source_url'))
    slide.notes_slide.notes_text_frame.text=json.dumps({'method':VERSION,'sources':rows,
        'images':image_sources,'image_policy':'Official web assets; exact case or clearly labelled similar operation / general tourism reference. No generated fallback.'},ensure_ascii=False,indent=2)


def selection(slide, report):
    header(slide,'이 사업을 참고한 이유와 지역 적용 방법')
    rows, primary=report_cases(report)
    source=primary[0] if primary else {}
    name=' · '.join(filter(None,[source.get('case_region'),source.get('intervention')]))
    text(slide,'selection-case',name or '동일 운영 방식의 공식 사례를 추가 연결할 수 있습니다.',106,198,1388,70,30,BLUE,True)
    facts=[f"{f.get('metric')}: {f.get('value')}" for f in report.get('observed_findings') or [] if f.get('metric') and f.get('value')]
    local=' / '.join(facts[:3]) or '저장된 지역 관측값이 없습니다.'
    strategy=base._first_strategy(report)
    comparison=[['비교 기준','확인한 근거','이번 기획에 연결한 판단'],
        ['사업의 이용 흐름',case_narrative(source.get('operating_model')) or '동일 방식 문서 미확보','결제·혜택·재이용 또는 예약·체류 등 실제 이용 흐름이 같은 사례만 연결'],
        ['선택 지역의 소비·체류',local,'지역 소비·체류 지표를 보고 참여 업종과 이용 대상을 정합니다.'],
        ['시범 운영 규모',(f"추가 방문 목표 {report['reference_estimate'].get('additional_visitors_target',0):,.0f}명 × 1%\n시범 참여 {report['reference_estimate']['quantity']:,}건(올림·계획 가정)" if report['reference_estimate'].get('additional_visitors_target') and 'quantity' in report['reference_estimate'] else '입력 예산 총액 안에서 참여량과 운영 인력을 배분'), '참여량 × 건별 혜택과 운영 인일로 견적을 계산합니다.']]
    table(slide,'selection-matrix',comparison,106,295,1388,360,[232,578,578],21)
    rect(slide,'adaptation-panel',106,680,1388,155,PALE)
    text(slide,'adaptation-label','사례 → 우리 지역의 제안',128,703,320,65,23,BLUE,True)
    safe_text(slide,'adaptation-body',strategy.get('solution') or strategy.get('title'),478,697,962,124,21,INK)
    label=text(slide,'selection-source','문서 기반 함수 연결 · '+str(source.get('title') or '기록 없음'),106,848,1388,38,18,BLUE)
    link(label,source.get('source_url'))
    # Details remain in speaker notes, outside the presentation narrative.
    # text(slide,'selection-limit','선정 근거: 운영 방식 일치. 인구·자연환경·교통이 같은 지역이라는 뜻이 아니며, 환경 유사성을 입증한 비교는 아닙니다.',106,862,1388,30,16,SLATE)
    slide.notes_slide.notes_text_frame.text=json.dumps({'method':VERSION,'selected_source':source,'observed_findings':report.get('observed_findings'),
        'interpretation':'7개 예측값으로 학습한 추천 모델 없음. 관측 지표는 지역 설계 입력이고 공식 운영 문서가 사례 연결 기준.'},ensure_ascii=False,indent=2)


def merged_execution(slide,report):
    header(slide,'4단계 실행 가이드 예시안')
    groups=execution_groups(report)
    for i,(lead,body) in enumerate(execution_copy(report)):
        y=220+i*143
        rect(slide,f'execution-panel-{i}',106,y,1388,126,WHITE)
        text(slide,f'flow-node-{i}',f'{i+1:02d}',128,y+22,75,51,34,CYAN,True)
        text(slide,f'execution-stage-{i}',LABELS[i],232,y+13,286,41,25,BLUE,True)
        text(slide,f'execution-period-{i}',period(groups[i])+' · '+ROLES[i],232,y+62,286,50,18,SLATE)
        text(slide,f'execution-lead-{i}',lead,558,y+10,906,37,23,INK,True)
        text(slide,f'execution-body-{i}',body,558,y+54,906,64,20,SLATE)
    text(slide,'execution-origin','작성 기준: 저장된 일정 + 사업 유형별 운영 안내 함수. 공식 협약·법정 절차를 인용한 내용이 아니라 조정 가능한 실행 예시입니다.',106,822,1388,60,19,SLATE)
    slide.notes_slide.notes_text_frame.text='proposal_layout_v9.execution_groups / execution_copy → 4개 단계. 일정은 저장 implementation_steps, 업무 설명은 환급형/일반형 코드 템플릿.\n'+json.dumps(base._first_strategy(report).get('implementation_steps'),ensure_ascii=False)


def estimate(slide,report):
    header(slide,'견적 예시안')
    e=report['reference_estimate']
    text(slide,'estimate-label','운영 규모를 가정한 시범 예산',106,203,890,50,29,SLATE,True)
    text(slide,'estimate-total',f"총 {e['total_krw']/10000:,.0f}만 원",1010,197,484,62,39,BLUE,True)

    rows=[['항목','금액','산출근거']]
    for i,item in enumerate(e['items']):
        rows.append([item['name'],f"{item['amount']:,}원",item['basis']])
    sh=table(slide,'estimate-table',rows,106,291,1388,402,[290,280,818],21)
    sh.table.cell(0,2).fill.solid();sh.table.cell(0,2).fill.fore_color.rgb=RGBColor.from_string('16845B')
    note=e.get('scale_basis', '참여량은 운영 기간과 예산에 따른 계획 가정입니다.')+'\n직접비 = 혜택 + 운영·정산 + 시스템 + 홍보 + 평가; 예비비 = 직접비 × 10%. 단가와 처리량은 기획 가정입니다.'
    if 'quantity' not in e:
        note='현재 금액은 사용자가 지정한 총액을 기존 항목 비중으로 재배분한 가정입니다. 각 항목 수량·단가는 별도 조정합니다.'
    text(slide,'estimate-method',note,106,720,1388,106,19,SLATE)
    text(slide,'estimate-note','예상 견적이며 실제 액수와 다를 수 있습니다. 세금 포함 여부·업무 범위·수량·단가는 계약 전에 조정합니다.',106,848,1388,38,18,BLUE)
    slide.notes_slide.notes_text_frame.text='proposal_evidence.build_reference_estimate; idea_proposal.scale_estimate.\n'+json.dumps(e,ensure_ascii=False,indent=2)


def polish(prs,report):
    project_targets(prs.slides[2], report)
    case_cards(prs.slides[4],report)
    selection(prs.slides[5],report)
    merged_execution(prs.slides[6],report)
    # Old slide 8 is fully consolidated into slide 7; no saved report is mutated.
    sid=prs.slides._sldIdLst[7];prs.part.drop_rel(sid.rId);prs.slides._sldIdLst.remove(sid)
    estimate(prs.slides[9],report)
    basis=prs.slides[8]
    method=next((s for s in basis.shapes if s.name=='basis-method'),None)
    if method:
        method.height=175*EMU
        data=report['target_proposal_basis'];actual=report['execution_scenario']
        explanation=('① 참고 실적: 강진 2024년 10월 누적 248만 명, 전년 대비 25% 증가. 반값여행·축제가 함께 운영된 지역 전체 실적.\n'
                     '② 목표 채택: 25% × 80% = 20%. 80%는 단계 운영을 고려한 계획 가정이며 인구·통계 보정값은 아님.\n'
                     if data.get('observed_visitors') else
                     '① 목표 채택: 공식 사례의 운영 방식을 참고해 초기 계획 목표 5% 제안.\n')
        explanation+=f"③ 적용 목표: 최종월 방문 +{actual['visitor_target_pct']:g}% · 소비 +{actual['spending_target_pct']:g}%. 월별 기준 전망에 단계 적용.\n"
        explanation+=('소비 목표는 해당 월 ML 소비/방문 비율을 추가 방문 목표에 적용한 계획값. 사례의 소비 증가율·사업 단독 효과가 아님.' if actual['visitor_target_pct']==actual['spending_target_pct'] else '소비 목표는 별도로 지정한 목표율을 해당 월 ML 소비액에 적용한 계획값. 사례의 소비 증가율·사업 단독 효과가 아님.')
        method.width=1348*EMU
        fit_text(method,explanation,23,SLATE)
    calculation_pages(prs, report, after=9)
    labels=['프로젝트 개요','사업 설계와 3개월 목표','지역별 참고 사례','사례 선정과 지역 적용','4단계 실행 가이드 예시안','ML 전망과 목표 KPI','사례 실적과 목표 설정','견적 예시안','산출 근거와 사례 선정 과정','근거·데이터·출처 및 맺음말']
    toc=prs.slides[1]
    for shape in list(toc.shapes):
        if shape.name.startswith('toc-'):shape._element.getparent().remove(shape._element)
    for i,label in enumerate(labels):
        col,row=divmod(i,5);x=562+540*col;y=296+94*row
        text(toc,f'toc-{i}-number',f'{i+1:02d}',x,y,61,44,25,CYAN,True)
        text(toc,f'toc-{i}-label',label,x+65,y+5,415,58,25,INK,True)


def project_targets(slide, report):
    from .proposal_calculation_basis import calculation_basis
    b=calculation_basis(report)
    for shape in list(slide.shapes):
        if shape.name in ('s3-operation', 's3-operation-rule'):
            shape._element.getparent().remove(shape._element)
    text(slide,'s3-goal-label','사업의 목표 · 3개월 월별 합계 기준',561,675,895,40,24,BLUE,True)
    for i,r in enumerate(b['totals']):
        x=561+i*459
        rect(slide,f's3-goal-panel-{i}',x,728,438,132,PALE)
        text(slide,f's3-goal-name-{i}',r['label'],x+15,739,400,30,22,BLUE,True)
        delta=(f"{r['additional']:,.0f}명" if r['key']=='visitors' else f"{r['additional']/1e8:,.2f}억 원")
        text(slide,f's3-goal-value-{i}',f'↑ {delta}',x+15,776,404,41,31,ORANGE,True)
        pct=f"+{r['pct']:.2f}%" if r['pct'] is not None else '비율 산정 제외'
        text(slide,f's3-goal-pct-{i}',f'기준 전망 대비 {pct} · 계획 목표',x+15,822,404,30,18,SLATE)
    if not b['totals']:
        text(slide,'s3-goal-missing','저장된 사업기간 전망과 목표율을 연결하면 방문·소비 목표를 표시합니다.',561,734,895,100,23,SLATE)


def calculation_pages(prs, report, after):
    from .proposal_calculation_basis import calculation_basis, methodology_sections
    from pptx.opc.packuri import PackURI
    # Earlier passes remove/reorder template slides. Renumber existing parts
    # before add_slide chooses len(slides)+1 to avoid duplicate ZIP part names.
    for index, slide in enumerate(prs.slides, 1):
        slide.part._partname=PackURI(f'/ppt/slides/slide{index}.xml')
    b=calculation_basis(report);sections=methodology_sections(report)
    from .proposal_infographics import calculations, ml_to_plan
    first=prs.slides.add_slide(prs.slides[-1].slide_layout)
    calculations(first,report)
    second=prs.slides.add_slide(prs.slides[-1].slide_layout)
    ml_to_plan(second,report)
    for slide in (first,second):
        script=(
            '발표 설명\n방문·소비의 기준 전망은 저장 모델에서 나온 수치입니다. 여기에 공식 사례를 참고한 계획 목표율을 단계적으로 적용합니다. '
            '강진을 참고한 경우 25%는 발표된 지역 증가율이고, 80%는 우리가 채택한 계획 가정입니다. 두 값을 곱한 20%는 신규 사업 효과의 ML 예측이 아닙니다. '
            '추가 방문 목표의 일부를 직접 운영할 참여량으로 배분하고 가정 단가를 곱합니다. 상세 항목의 합계에 예비비를 더한 값이 견적입니다.\n\n'
            if slide==first else
            '발표 설명\n우리 ML은 7개 관광지표의 전망을 계산합니다. 방문자와 소비액은 그래프·목표 규모를 계산하는 기준선입니다. '
            '체류·숙박·검색 관련 5개 전망은 지역의 관광 흐름을 검토할 입력입니다. 서버가 이 수치와 모델 평가를 공식 사례와 함께 Qwen에 제공합니다. '
            'Qwen은 지역 조건과 사례의 운영 방식을 비교하고 Gemma는 선택안을 본문으로 구성합니다. 따라서 ML은 수치 전망을 담당하고 사업 선택은 근거 문서와 LLM 판단을 결합합니다. '
            '지표별로 전년 동월 기준모델이 선택될 수 있으며, 이를 고도화 모델의 성능 개선이라고 설명하지 않습니다.\n\n')
        slide.notes_slide.notes_text_frame.text=script+json.dumps({'sections':sections,'basis':b,
            'functions':['planning_evidence.build_planning_ml_evidence','case_recommendation.link_decision',
                         'report_projection.execution_target','proposal_evidence.build_reference_estimate']},ensure_ascii=False,indent=2)
    # Insert the two new pages directly after the estimate, ahead of bibliography.
    for offset,slide in enumerate((first,second),1):
        sid=next(sid for sid in prs.slides._sldIdLst if sid.id==slide.slide_id)
        prs.slides._sldIdLst.remove(sid);prs.slides._sldIdLst.insert(after+offset,sid)
