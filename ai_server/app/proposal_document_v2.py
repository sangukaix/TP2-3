"""A4 proposal, sharing the approved PPT's data and narrative projections."""
from io import BytesIO
import re

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .idea_proposal import prepare_idea_report
from .proposal_calculation_basis import calculation_basis, methodology_sections
from .proposal_presentation_v4 import _scenario_for_display, _public_copy
from .proposal_evidence import complete_source_records, source_display_name
from .case_recommendation import report_cases, case_reference_role
from .case_images import case_image
from .proposal_layout_v9 import execution_groups, execution_copy, period, LABELS, ROLES

BLUE = '0054A6'
GREEN = '16845B'
GRAY = '536477'


def prose(value):
    value = re.sub(r'\s*\(?\[[^\]]+\]\(https?://[^\s]+\)\)?', '', _public_copy(value))
    return value.strip()


def paragraph(parent, value='', *, size=10, bold=False, color='263445', style=None):
    p = parent.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.22
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
    p._p.get_or_add_pPr().append(snap)
    r = p.add_run(str(value))
    r.font.name = '맑은 고딕'
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def heading(doc, value, new_page=False):
    p = paragraph(doc, value, size=19, bold=True, color='000000', style='Heading 1')
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.space_after = Pt(15)
    return p


def subheading(doc, value):
    p = paragraph(doc, value, size=12, bold=True, color='000000', style='Heading 2')
    p.paragraph_format.space_before = Pt(10)
    return p


def table(doc, rows, widths, green_last=False):
    t = doc.add_table(rows=0, cols=len(widths))
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Mm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement('w:' + edge)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            element.set(qn('w:' + key), val)
        borders.append(element)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = t.add_row()
        if i == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), (GREEN if green_last and j == len(widths)-1 else BLUE) if i == 0 else ('F0F5F8' if i % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '100'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            props.append(margins)
            p = paragraph(cell, value, size=9, bold=i == 0, color='FFFFFF' if i == 0 else '263445')
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.14
            cell._tc.remove(cell.paragraphs[0]._p)
    paragraph(doc, '', size=2).paragraph_format.space_after = Pt(2)
    return t


def forecast_chart(scenario, metric=None):
    # The Word plot consumes exactly the arrays that drive the native PPT charts.
    from .proposal_document import plt
    specs = [('visitors', '방문자 수 · 만 명', 10000), ('spending', '관광소비액 · 억 원', 100000000)]
    if metric: specs = [v for v in specs if v[0] == metric]
    fig, axes = plt.subplots(1, len(specs), figsize=(8.1, 3.05), dpi=190, squeeze=False)
    for ax, (key, label, unit) in zip(axes[0], specs):
        baseline = [v/unit for v in scenario['baseline_' + key]]
        ax.plot(range(len(baseline)), baseline, color='#0054A6', marker='o', lw=2.3, label='ML 기준 전망')
        if scenario.get('has_target'):
            ax.plot(range(len(baseline)), [v/unit for v in scenario['target_' + key]], color='#FF7400', marker='o', lw=2.3, label='목표 KPI')
        ax.set_title(label, loc='left', fontsize=11, pad=15)
        ax.set_xticks(range(len(baseline)), scenario['categories'], fontsize=8)
        ax.tick_params(axis='y', labelsize=8)
        ax.grid(axis='y', color='#DDE6EC'); ax.set_axisbelow(True)
        ax.margins(x=.14, y=.25)
        for side in ('top', 'right'): ax.spines[side].set_visible(False)
        ax.legend(loc='upper center', bbox_to_anchor=(.5, -.2), ncol=2, frameon=False, fontsize=8)
    fig.subplots_adjust(left=.07, right=.98, top=.83, bottom=.28, wspace=.32)
    output = BytesIO(); fig.savefig(output, format='png', facecolor='white'); plt.close(fig); output.seek(0)
    return output


def metric_rows(scenario, key):
    money = key == 'spending'
    fmt = (lambda v: f'{v/100000000:,.2f}억 원') if money else (lambda v: f'{v:,.0f}명')
    rows = [['기준월', 'ML 기준 전망', '목표 KPI', '기준 대비 증가']]
    base = scenario['baseline_' + key]
    targets = scenario.get('target_' + key) if scenario.get('has_target') else None
    for month, baseline, target in zip(scenario['categories'], base, targets or [None]*len(base)):
        diff = target-baseline if target is not None else None
        change = f'↑ {diff/baseline*100:.2f}%\n+{fmt(diff)}' if diff is not None and baseline > 0 else '—'
        rows.append([month, fmt(baseline), fmt(target) if target is not None else '—', change])
    if targets:
        b, v = sum(base), sum(targets)
        rows.append(['월별 합계', fmt(b), fmt(v), f'↑ {(v/b-1)*100:.2f}%\n+{fmt(v-b)}' if b else '—'])
    return rows


def create_strategy_proposal_document(report):
    report = prepare_idea_report(report)
    strategy = (report.get('strategies') or [{}])[0]
    b = calculation_basis(report)
    scenario, demo = _scenario_for_display(report)
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(19); sec.bottom_margin = Mm(18)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.header_distance = sec.footer_distance = Mm(9)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]; style.font.name = '맑은 고딕'; style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        for border in list(style._element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(10)
    header = sec.header.paragraphs[0]
    header.text = f"{report.get('region_name', '')}  |  관광 전략기획안"
    header.runs[0].font.size = Pt(8); header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('OLIGO-K   ·   ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    paragraph(doc, report.get('region_name', ''), size=11, bold=True, color=BLUE)
    paragraph(doc, prose(strategy.get('title') or '관광 전략기획안'), size=25, bold=True, color='000000', style='Title')
    paragraph(doc, '관광 전략기획안', size=13, color=GRAY)
    paragraph(doc, '사업기간  '+(' ~ '.join([b['months'][0], b['months'][-1]]) if b['months'] else str(strategy.get('timeframe') or '')), bold=True)
    subheading(doc, '핵심 제안')
    paragraph(doc, prose(strategy.get('summary') or report.get('summary')))
    subheading(doc, '사업의 목표')
    rows = [['지표', '3개월 목표 합계', '기준 전망 대비 추가 규모']]
    for r in b['totals']:
        fmt = (lambda v: f'{v/1e8:,.2f}억 원') if r['key']=='spending' else (lambda v: f'{v:,.0f}명')
        rows.append([r['label'], fmt(r['target']), f"↑ {r['pct']:.2f}% · +{fmt(r['additional'])}" if r['pct'] is not None else fmt(r['additional'])])
    if len(rows)>1: table(doc, rows, [38, 60, 72])
    if demo: paragraph(doc, '예시 데이터로 만든 시연용 목표입니다.', color=GRAY)
    paragraph(doc, '목표는 월별 ML 기준 전망에 계획 증가율을 적용한 값입니다. 방문 합계는 월간 순 방문 수의 합계이며, 3개월 동안 중복을 제거한 인원이 아닙니다.', size=9, color=GRAY)
    subheading(doc, '제안의 방향')
    paragraph(doc, prose(strategy.get('problem_to_solve') or strategy.get('problem')))
    paragraph(doc, prose(strategy.get('solution')))
    subheading(doc, '문서 구성')
    paragraph(doc, '01 전망과 목표  ·  02 지역별 참고 사례  ·  03 지역 적용\n04 실행 가이드  ·  05 견적  ·  06 산출 근거  ·  07 데이터 활용  ·  출처', size=9, color=GRAY)

    heading(doc, '01  머신러닝 전망과 목표 KPI', True)
    if scenario:
        for key, name in [('visitors', '방문자 수'), ('spending', '관광소비액')]:
            if key == 'spending': heading(doc, '01  머신러닝 전망과 목표 KPI 계속', True)
            subheading(doc, name)
            doc.add_picture(forecast_chart(scenario, key), width=Mm(170))
            table(doc, metric_rows(scenario, key), [25, 47, 47, 51])
            subheading(doc, '수치 해석과 계산 기준')
            paragraph(doc, methodology_sections(report)[0][1], size=9)
            paragraph(doc, '월별 목표 = 월별 기준 전망 × (1 + 최종월 목표율 × 경과월/전체월). 합계 증가율은 월별 목표 합계 ÷ 기준 전망 합계 − 1로 계산합니다.', size=9)
            paragraph(doc, '소비액 표시는 억 원 단위 소수 둘째 자리로 반올림합니다. 합계와 증가율은 반올림 전 값으로 계산합니다. 방문 합계는 3개월 고유 인원이 아닌 월간 순 방문 수의 합계입니다.', size=8, color=GRAY)
    else:
        paragraph(doc, '해당 사업기간에 저장된 전망 수치가 없습니다. 관측값을 예측 그래프로 대체하지 않습니다.')

    heading(doc, '02  지역별 참고 사례', True)
    paragraph(doc, '공식 문서에서 사업의 실제 운영 방식이 연결되는 사례를 참고합니다. 아래 구분은 선정 사업의 근거와 추가 참고 사례의 역할을 나타냅니다.', size=9, color=GRAY)
    cases, primary = report_cases(report)
    for i, source in enumerate(cases):
        start = len(doc.paragraphs)
        subheading(doc, f"{i+1:02d}  {source.get('case_region') or ''}  {source.get('intervention') or source.get('title') or ''}")
        paragraph(doc, case_reference_role(report, source), size=9, bold=True, color=BLUE)
        paragraph(doc, prose(source.get('operating_model')), size=9)
        info = case_image(source)
        if info:
            from PIL import Image
            with Image.open(info['path']) as im:
                ratio = min(150/im.width, 25/im.height)
                doc.add_picture(str(info['path']), width=Mm(im.width*ratio), height=Mm(im.height*ratio))
            paragraph(doc, info['caption']+' · '+info['credit'], size=8, color=GRAY)
            paragraph(doc, info['page_url'], size=7, color=GRAY)
        paragraph(doc, '사업 원문  '+str(source.get('source_url') or source.get('source_id') or ''), size=7, color=GRAY)
        for p in doc.paragraphs[start:-1]:
            p.paragraph_format.keep_with_next = True

    heading(doc, '03  사례 선정과 지역 적용', True)
    source = primary[0] if primary else {}
    subheading(doc, '사업의 이용 흐름')
    paragraph(doc, prose(source.get('operating_model')))
    paragraph(doc, '결제·혜택·재이용 또는 예약·체류 등 실제 이용 흐름이 같은 사례를 연결합니다.', size=9, color=GRAY)
    subheading(doc, '선택 지역의 소비와 체류')
    facts = [(f.get('metric'), f.get('value')) for f in report.get('observed_findings') or [] if f.get('metric') and f.get('value')]
    if facts: table(doc, [['관측 지표', '확인한 값']]+facts[:3], [70, 100])
    paragraph(doc, '지역 소비·체류 지표를 보고 참여 업종과 이용 대상을 정합니다.')
    subheading(doc, '운영 규모')
    paragraph(doc, b['estimate'].get('scale_basis') or '입력 예산 총액 안에서 참여량과 운영 인력을 배분합니다.')
    subheading(doc, '사례에서 우리 지역의 제안으로')
    paragraph(doc, prose(strategy.get('solution')))
    if strategy.get('comparison_analysis'):
        subheading(doc, '선정 판단')
        paragraph(doc, prose(strategy['comparison_analysis']), size=9)
    paragraph(doc, '연결 근거  '+str(source.get('title') or ''), size=8, color=GRAY)

    heading(doc, '04  4단계 실행 가이드 예시안', True)
    groups = execution_groups(report)
    for i, (lead, body) in enumerate(execution_copy(report)):
        subheading(doc, f'{i+1:02d}  {LABELS[i]}')
        paragraph(doc, period(groups[i])+'  |  '+ROLES[i], size=9, color=BLUE)
        paragraph(doc, lead, bold=True)
        paragraph(doc, body)
    paragraph(doc, '일정은 저장 기획안, 업무 설명은 사업 유형별 운영 안내 함수를 사용한 실행 예시입니다. 담당 역할과 운영 조건은 협의해 조정할 수 있습니다.', size=9, color=GRAY)

    heading(doc, '05  견적 예시안', True)
    e = b['estimate']
    paragraph(doc, f"총 {e.get('total_krw', 0):,}원", size=24, bold=True, color=BLUE)
    paragraph(doc, '예상 견적이며 실제 액수와 다를 수 있습니다.', color=GRAY)
    table(doc, [['항목', '금액', '산출근거']]+[[v['name'], f"{v['amount']:,}원", v['basis']] for v in e.get('items') or []]+[['합계', f"{e.get('total_krw', 0):,}원", '항목별 금액 합계']], [36, 36, 98], green_last=True)
    subheading(doc, '규모를 정한 기준')
    paragraph(doc, e.get('scale_basis') or '사용자 지정 예산 총액을 항목별로 배분한 계획 가정입니다.')
    paragraph(doc, '직접비 = 참여 혜택 + 운영·정산 + 시스템 + 홍보 + 평가. 예비비 = 직접비 × 10%. 단가와 처리량은 기획 가정입니다.', size=9)
    for note in e.get('assumptions') or []: paragraph(doc, note, size=8, color=GRAY)

    heading(doc, '06  전망에서 목표와 견적까지', True)
    methods = methodology_sections(report)
    subheading(doc, '1  월별 기준 전망을 계산합니다')
    paragraph(doc, methods[0][1])
    subheading(doc, '2  사례 실적을 목표율로 해석합니다')
    paragraph(doc, b['target_explanation'], size=9)
    subheading(doc, '3  월별 목표와 추가 규모를 계산합니다')
    paragraph(doc, methods[1][1], size=9)
    subheading(doc, '4  참여량과 운영비를 편성합니다')
    paragraph(doc, e.get('scale_basis') or '지정된 예산 총액을 항목별로 배분합니다.', size=9)
    paragraph(doc, '참여량·단가·운영 처리량은 기획 가정입니다. 운영 예산이 지역 전체 목표 증가를 발생시킨다는 인과 계산은 아닙니다.', size=9, color=GRAY)

    heading(doc, '07  머신러닝 결과가 기획에 쓰이는 과정', True)
    table(doc, [['입력', '결과의 사용처'], ['방문자 수 · 관광소비액 전망', '전망 그래프 → 월별 목표 → 추가 규모 → 견적의 운영량 가정'], ['체류시간 · 숙박일수 · 숙박방문 비율\n내비게이션 검색 · 숙박 검색 전망', '관광 흐름 진단 → 공식 사례 조사 방향 → 후보별 지역 적용 검토']], [75, 95])
    subheading(doc, '공식 근거에서 최종 기획까지')
    paragraph(doc, f"저장된 공식 사례 {b['case_count']}건 → 설계 후보 {b['candidate_count']}개 → 참고 사례 {b['displayed_count']}건 표시", size=12, bold=True, color=BLUE)
    for title, body in [
        ('01  지역 데이터와 전망', '원자료의 관측값과 7개 지표 전망을 지역 기획 입력으로 구성합니다. 방문·소비는 수치 기준선이며 나머지는 조사와 설계 방향을 검토하는 진단 입력입니다.'),
        ('02  Qwen의 후보 비교', '공식 사례의 운영 방식과 지역 입력을 읽고 후보를 비교합니다. 함수가 사업 운영 방식과 맞는 공식 문서를 연결합니다. 7개 예측값으로 학습한 지역 추천 모델은 아닙니다.'),
        ('03  Gemma의 본문 작성', '선정 후보와 지역 수치·공식 근거를 바탕으로 사업 소개와 실행 내용을 작성합니다. 목표와 견적은 별도의 계획 가정과 Python 산식으로 구성합니다.'),
        ('04  검수와 문서 출력', '수치·출처·기간을 확인한 저장 보고서에서 PPT와 Word를 생성합니다. 문서 출력 과정에서 LLM을 새로 호출하거나 예측 모델을 다시 학습하지 않습니다.')]:
        subheading(doc, title); paragraph(doc, body)

    heading(doc, '참고 자료와 출처', True)
    paragraph(doc, '저장 보고서의 전체 근거 목록입니다. 원자료·ML 전망·공식 사례를 구분하고 원문 주소와 문서 위치를 유지합니다.', size=9, color=GRAY)
    for i, source in enumerate(complete_source_records(report), 1):
        p = paragraph(doc, f"{i:02d}  {source_display_name(source)}", size=9, bold=True)
        p.paragraph_format.keep_with_next = True
        details = [str(source.get('source_type') or ''), str(source.get('source_id') or '')]
        if source.get('page_references'): details.append('페이지 '+str(source['page_references']))
        if source.get('chunk_ids'): details.append('문서 구간 '+', '.join(source['chunk_ids']))
        p.paragraph_format.space_after = Pt(2)
        p = paragraph(doc, ' · '.join(filter(None, details)), size=8, color=GRAY)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = bool(source.get('source_url'))
        if source.get('source_url'):
            paragraph(doc, source['source_url'], size=8, color=BLUE)
    out = BytesIO(); doc.save(out); out.seek(0)
    return out
