"""사업 여건 검증·Agent 전달·작업 고정·문서 추출을 유료 API 없이 검증합니다."""
import asyncio
from io import BytesIO
import unittest
from unittest.mock import AsyncMock, patch

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from fastapi.testclient import TestClient
from pydantic import ValidationError

from ai_server.app import main
from ai_server.app.planning_brief import PlanningBrief, brief_fingerprint, extract_brief_reference
from ai_server.app.agents.evidence_agent import EvidenceAgent
from ai_server.app.agents.case_study_agent import CaseStudyAgent
from ai_server.app.agents.transferability_agent import TransferabilityAgent
from ai_server.app.agents.chat_assistant_agent import TourismChatAssistantAgent
from ai_server.app.agents.report_orchestrator import _EVIDENCE_CACHE, _CASE_STUDY_CACHE, orchestrate_strategy_report
from ai_server.test_report_orchestrator import FakeEvidenceAgent, FakeCaseStudyAgent, FakePlannerAgent, FakeReviewerAgent, FakeTransferabilityAgent
from pathlib import Path
from zipfile import ZipFile


def brief(**changes):
    return PlanningBrief(region_code='11680', **changes)


class BriefContractTest(unittest.TestCase):
    def test_unknown_remains_unknown(self):
        value = brief().model_dump(mode='json')
        self.assertIsNone(value['budget_max_krw'])
        self.assertIsNone(value['start_date'])
        self.assertEqual(value['resources_confirmed'], '')

    def test_invalid_conditions_rejected(self):
        for changes in [
            {'budget_status': 'confirmed'},
            {'budget_status': 'unknown', 'budget_max_krw': 100},
            {'budget_status': 'indicative', 'budget_max_krw': 100, 'budget_min_krw': 200},
            {'budget_status': 'confirmed', 'budget_max_krw': True},
            {'schedule_status': 'fixed', 'start_date': '2026-12-01', 'end_date': '2026-11-01'},
            {'schedule_status': 'fixed'},
            {'resources_confirmed': 'x' * 1501},
        ]:
            with self.subTest(changes=changes), self.assertRaises(ValidationError):
                brief(**changes)

    def test_hash_tracks_conditions(self):
        a = brief().model_dump(mode='json')
        b = brief(hard_constraints='야간 운영 제외').model_dump(mode='json')
        self.assertNotEqual(brief_fingerprint(a), brief_fingerprint(b))
        self.assertEqual(brief_fingerprint(a), brief_fingerprint(dict(reversed(list(a.items())))))

    def test_text_docx_pdf_and_xlsx_extraction(self):
        self.assertEqual(extract_brief_reference('memo.txt', '현장 정보'.encode())['text'], '현장 정보')
        doc = Document(); doc.add_paragraph('사용자 제공 참고자료')
        buffer = BytesIO(); doc.save(buffer)
        result = extract_brief_reference('../memo.docx', buffer.getvalue())
        self.assertEqual(result['name'], 'memo.docx')
        self.assertIn('참고자료', result['text'])
        pdf = PdfWriter(); pdf.add_blank_page(width=72, height=72); pdf_buffer = BytesIO(); pdf.write(pdf_buffer)
        # 빈 PDF는 첨부할 수 없으므로, PDF 형식 오류가 아닌 텍스트 부재로 거절됩니다.
        with self.assertRaises(ValueError):
            extract_brief_reference('memo.pdf', pdf_buffer.getvalue())
        workbook = Workbook(); workbook.active.title = '예산'; workbook.active.append(['항목', '금액']); workbook.active.append(['홍보', 3000000]); xlsx_buffer = BytesIO(); workbook.save(xlsx_buffer)
        self.assertIn('홍보', extract_brief_reference('budget.xlsx', xlsx_buffer.getvalue())['text'])
        hwpx_buffer = BytesIO()
        with ZipFile(hwpx_buffer, 'w') as archive:
            archive.writestr('Contents/section0.xml', '<hp:section xmlns:hp="urn:hancom">현장 의견</hp:section>')
        self.assertIn('현장 의견', extract_brief_reference('memo.hwpx', hwpx_buffer.getvalue())['text'])

    def test_invalid_attachments_rejected(self):
        for filename, content in [('memo.pdf', b'pdf'), ('memo.hwp', b'legacy-hwp'), ('memo.txt', b'a'*6001), ('memo.txt', b'a'*2_000_001), ('memo.docx', b'not-a-zip')]:
            with self.subTest(filename=filename), self.assertRaises(ValueError):
                extract_brief_reference(filename, content)

    def test_api_validation_no_llm(self):
        with TestClient(main.app) as client, patch('ai_server.app.main.generate_orchestrated_report', new_callable=AsyncMock) as generate:
            response = client.post('/ai/v1/demo/11680/strategy-report/jobs', json={'region_name': '서울특별시 강남구', 'planning_brief': {'region_code': '11680', 'budget_status': 'confirmed'}})
            self.assertEqual(response.status_code, 422)
            response = client.post('/ai/v1/demo/11680/strategy-report/jobs', json={'region_name': '서울특별시 강남구', 'planning_brief': {'region_code': '28245'}})
            self.assertEqual(response.status_code, 422)
            response = client.post('/ai/v1/planning/reference?filename=note.txt', content='현장 메모'.encode(), headers={'Content-Type': 'application/octet-stream'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()['text'], '현장 메모')
            generate.assert_not_called()


class BriefAgentFlowTest(unittest.IsolatedAsyncioTestCase):
    async def test_changed_brief_invalidates_research_cache(self):
        _EVIDENCE_CACHE.clear(); _CASE_STUDY_CACHE.clear()
        FakeEvidenceAgent.calls = 0; FakeCaseStudyAgent.calls = 0; FakeReviewerAgent.calls = 1
        settings = {'project_root': Path('.'), 'env_values': {'OPENAI_API_KEY': 'test'},
                    'region_code': '11680', 'snapshot': {'region_name': '서울특별시 강남구'}, 'report_schema': {}}
        with (
            patch('ai_server.app.agents.report_orchestrator.EvidenceAgent', FakeEvidenceAgent),
            patch('ai_server.app.agents.report_orchestrator.CaseStudyAgent', FakeCaseStudyAgent),
            patch('ai_server.app.agents.report_orchestrator.TransferabilityAgent', FakeTransferabilityAgent),
            patch('ai_server.app.agents.report_orchestrator.PlannerAgent', FakePlannerAgent),
            patch('ai_server.app.agents.report_orchestrator.ReviewerAgent', FakeReviewerAgent),
        ):
            a = brief(hard_constraints='야간 제외').model_dump(mode='json')
            b = brief(hard_constraints='신규 시설 제외').model_dump(mode='json')
            await orchestrate_strategy_report(**settings, planning_brief=a)
            await orchestrate_strategy_report(**settings, planning_brief=a)
            await orchestrate_strategy_report(**settings, planning_brief=b)
            self.assertEqual(FakeEvidenceAgent.calls, 2)
            self.assertEqual(FakeCaseStudyAgent.calls, 2)
            self.assertEqual(FakePlannerAgent.last_evidence_pack['planning_brief'], b)

    async def test_direct_agents_receive_separate_conditions(self):
        conditions = brief(hard_constraints='야간 제외').model_dump(mode='json')
        snapshot = {'region_name': '서울특별시 강남구', 'period': '2026-07'}
        empty = AsyncMock(return_value=[])
        with (
            patch('ai_server.app.agents.evidence_agent.OfficialTourismRagStore.search', empty),
            patch('ai_server.app.agents.evidence_agent.TourismOpenApiClient.collect_region_resources', empty),
            patch('ai_server.app.agents.evidence_agent.create_structured_response', AsyncMock(return_value={'summary': '', 'findings': [], 'gaps': []})) as request,
        ):
            await EvidenceAgent(project_root=Path('.'), env_values={'OPENAI_API_KEY': 'test'}).collect(region_code='11680', snapshot=snapshot, planning_brief=conditions)
            self.assertEqual(request.call_args.kwargs['input_payload']['planning_brief'], conditions)
        with (
            patch('ai_server.app.agents.case_study_agent.OfficialTourismRagStore.search', empty),
            patch('ai_server.app.agents.case_study_agent.create_structured_response', AsyncMock(return_value={'summary': '', 'cases': [], 'gaps': []})) as request,
        ):
            await CaseStudyAgent(project_root=Path('.'), env_values={'OPENAI_API_KEY': 'test'}).collect(region_code='11680', snapshot=snapshot, planning_brief=conditions)
            self.assertEqual(request.call_args.kwargs['input_payload']['planning_brief'], conditions)
        with patch('ai_server.app.agents.transferability_agent.create_structured_response', AsyncMock(return_value={})) as request:
            await TransferabilityAgent(api_key='test', model='test').assess(evidence_pack={'snapshot': snapshot, 'benchmark_cases': [{}], 'planning_brief': conditions})
            self.assertEqual(request.call_args.kwargs['input_payload']['planning_brief'], conditions)
        with patch('ai_server.app.agents.chat_assistant_agent.create_structured_response', AsyncMock(return_value={'sources': []})) as request:
            await TourismChatAssistantAgent(env_values={'OPENAI_API_KEY': 'test'}).answer(snapshot=snapshot, question='질문', history=[], current_report=None, enable_web_search=False, planning_brief=conditions)
            self.assertEqual(request.call_args.kwargs['input_payload']['planning_brief'], conditions)
        self.assertNotIn('planning_brief', snapshot)

    async def test_response_retains_frozen_brief(self):
        conditions = brief(budget_status='confirmed', budget_max_krw=30_000_000, references=[{'name': 'memo.txt', 'text': '첨부 본문'}])
        request = main.ReportRequest(region_name='서울특별시 강남구', planning_brief=conditions)
        generated = {'report': {'summary': '테스트', 'observed_findings': [], 'strategies': []},
                     'quality_review': {}, 'evidence_sources': [], 'research_gaps': [], 'agent_trace': []}
        with (
            patch.dict(main.ENV_VALUES, {'OPENAI_API_KEY': 'test'}),
            patch('ai_server.app.main.build_region_snapshot', return_value={'region_name': request.region_name, 'period': '2026-07', 'observations': [], 'monthly_trend': []}),
            patch('ai_server.app.main.orchestrate_strategy_report', AsyncMock(return_value=generated)) as orchestrator,
        ):
            result = await main.generate_orchestrated_report('11680', request)
            self.assertEqual(orchestrator.call_args.kwargs['planning_brief']['budget_max_krw'], 30_000_000)
            self.assertEqual(result.planning_brief.budget_max_krw, 30_000_000)
            self.assertEqual(result.planning_brief.references, [])
            self.assertTrue(result.planning_brief_fingerprint)


if __name__ == '__main__':
    unittest.main()
