"""Word/PPT numerical parity and data retention, without network or LLM."""
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
import unittest

from docx import Document
from ai_server.test_proposal_presentation_v3_contract import _sample_report
from ai_server.app.idea_proposal import prepare_idea_report
from ai_server.app.proposal_document import create_strategy_proposal_document
from ai_server.app.proposal_document_v2 import forecast_chart, metric_rows
from ai_server.app.proposal_presentation_v4 import _scenario_for_display


class WordParityTest(unittest.TestCase):
    def test_plot_uses_ppt_baseline_and_goal_arrays_without_recalculation(self):
        from matplotlib.axes import Axes
        r = prepare_idea_report(_sample_report())
        scenario, _ = _scenario_for_display(r)
        captured = []
        original = Axes.plot
        def capture(axis, x, y, **kwargs):
            captured.append((list(y), kwargs['label']))
            return original(axis, x, y, **kwargs)
        with patch.object(Axes, 'plot', capture):
            result = forecast_chart(scenario)
        self.assertTrue(result.getvalue().startswith(b'\x89PNG'))
        self.assertEqual(captured, [
            ([v/10000 for v in scenario['baseline_visitors']], 'ML 기준 전망'),
            ([v/10000 for v in scenario['target_visitors']], '목표 KPI'),
            ([v/100000000 for v in scenario['baseline_spending']], 'ML 기준 전망'),
            ([v/100000000 for v in scenario['target_spending']], '목표 KPI'),
        ])
        captured.clear()
        with patch.object(Axes, 'plot', capture):
            forecast_chart(scenario, 'spending')
        self.assertEqual(len(captured), 2)
        self.assertEqual(captured[0][0], [v/100000000 for v in scenario['baseline_spending']])
        self.assertEqual(captured[1][0], [v/100000000 for v in scenario['target_spending']])

    def test_document_retains_sources_and_exact_budget_without_mutating_report(self):
        r = _sample_report()
        long_title = '공식 원문 자료 제목 '*24 + '마지막 원문까지 유지'
        r['evidence_sources'].append({'source_id':'dataset:full-tail', 'source_type':'dataset',
            'title':long_title, 'source_url':'https://example.go.kr/full-source', 'page_references':'12, 19'})
        before = deepcopy(r)
        doc = Document(create_strategy_proposal_document(r))
        all_text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertEqual(r, before)
        self.assertIn(long_title, all_text)
        self.assertIn('https://example.go.kr/full-source', all_text)
        self.assertIn('페이지 12, 19', all_text)
        self.assertIn('4단계 실행 가이드', all_text)
        self.assertNotIn('5단계 집행 방법', all_text)
        self.assertAlmostEqual(doc.sections[0].page_width.mm, 210, delta=.1)
        e = prepare_idea_report(r)['reference_estimate']
        for item in e['items']: self.assertIn(f"{item['amount']:,}원", all_text)
        self.assertIn(f"{e['total_krw']:,}원", all_text)

    def test_monthly_table_matches_shared_values_and_has_no_observation_fallback(self):
        r = prepare_idea_report(_sample_report())
        scenario, _ = _scenario_for_display(r)
        rows = metric_rows(scenario, 'visitors')
        self.assertEqual(len(rows), len(scenario['categories'])+2)
        for i, row in enumerate(rows[1:-1]):
            self.assertEqual(row[0], scenario['categories'][i])
            self.assertEqual(row[1], f"{scenario['baseline_visitors'][i]:,.0f}명")
            self.assertEqual(row[2], f"{scenario['target_visitors'][i]:,.0f}명")
        no_ml = _sample_report(with_ml=False)
        doc = Document(create_strategy_proposal_document(no_ml))
        self.assertIn('관측값을 예측 그래프로 대체하지 않습니다.', '\n'.join(p.text for p in doc.paragraphs))


if __name__ == '__main__': unittest.main()
