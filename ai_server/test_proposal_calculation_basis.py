import unittest
from zipfile import ZipFile
from pptx import Presentation
from ai_server.test_proposal_presentation_v3_contract import _sample_report
from ai_server.app.proposal_presentation_v4 import create_strategy_proposal_presentation
from ai_server.app.proposal_layout_v7 import project, EMU, BLUE
from pptx.enum.text import MSO_ANCHOR


class CalculationPagesTest(unittest.TestCase):
    def test_caption_grows_and_stays_vertically_centered(self):
        from ai_server.app.proposal_layout_v7 import text,rect
        report=_sample_report(); deck=Presentation()
        for _ in range(3): deck.slides.add_slide(deck.slide_layouts[6])
        slide=deck.slides[2]
        label=text(slide,'hero-source-label','',30,800,450,50,17)
        background=rect(slide,'hero-source-bg',30,800,450,50,BLUE)
        project(deck,report,[{'address':'제주시','title':'가마오름'}])
        one=label.height
        self.assertEqual(label.text_frame.vertical_anchor,MSO_ANCHOR.MIDDLE)
        self.assertEqual(label.top+label.height//2,background.top+background.height//2)
        project(deck,report,[{'address':'제주특별자치도 제주시 한경면 청수리 상세 주소와 위치','title':'가마오름 관광 안내 사진'}])
        self.assertGreater(label.height,one)
        self.assertEqual(label.top+label.height//2,background.top+background.height//2)

    def test_new_pages_are_after_estimate_and_have_unique_package_parts(self):
        output = create_strategy_proposal_presentation(_sample_report())
        with ZipFile(output) as archive:
            self.assertEqual(len(archive.namelist()), len(set(archive.namelist())))
        output.seek(0)
        deck = Presentation(output)
        title = lambda i: '\n'.join(s.text for s in deck.slides[i].shapes if s.has_text_frame)
        self.assertIn('견적 예시안', title(9))
        self.assertIn('산출 근거 ①', title(10))
        self.assertIn('산출 근거 ②', title(11))
        self.assertNotIn('s3-operation', [s.name for s in deck.slides[2].shapes])
        self.assertIn('사업의 목표', title(2))
