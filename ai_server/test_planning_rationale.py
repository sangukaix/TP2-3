"""Provenance and backward compatibility, without LLM or database calls."""
from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import AsyncMock
import unittest

from jsonschema import Draft202012Validator, ValidationError
from ai_server.app.planning_rationale import (
    fact_catalog, bind_rationale, report_rationale, rationale_sections,
)
from ai_server.app.agents.transferability_agent import TransferabilityAgent
from ai_server.app.llm.evidence_tools import EvidenceTools


def pack():
    return {'sources': [{'source_id': 'dataset:1:1', 'source_type': 'dataset',
                        'summary': '숙박 방문 비율 6.6%', 'observation_period': '2026-06'}],
            'snapshot': {'ml_analysis': {'signals': [
                {'source_id': 'ml:1:v1', 'metric': 'visitors', 'period': '202610~202612',
                 'aggregation': '3_month_sum', 'forecast_value': 120, 'unit': '명',
                 'model_reliability': 'below_baseline_on_test'}]}},
            'benchmark_cases': [{'source_id': 'case:stamp', 'source_url': 'https://example.go.kr/case',
                                 'operating_model': 'GPS 스탬프 투어 완주 쿠폰 지급'}]}


def decision():
    return {'selected_candidate_id': 'stamp', 'selection_status': 'needs_evidence',
            'selection_reason': '완주 후 시장 재방문을 검토하되 야간 체류 연장은 별도 대안이다.',
            'design_candidates': [{'candidate_id': 'stamp', 'title': '스탬프 코스',
                'mechanism': 'GPS 스탬프 완주 쿠폰', 'case_source_ids': ['case:stamp'],
                'evidence_source_ids': [], 'reasoning': {
                    'fact_ids': list(fact_catalog(pack())),
                    'application_hypothesis': '완주 후 시장 방문으로 연결해 볼 수 있다.',
                    'visitor_action': '완주 후 쿠폰을 사용한다.',
                    'operating_requirements': '가맹점·정산 협약을 준비한다.',
                    'measurement': '지급 ID 중 사용 ID를 집계한다.',
                    'tradeoff': '가맹점 참여는 아직 확정 전이다.'}}]}


class RationaleTest(unittest.TestCase):
    def test_ppt_adds_trace_page_only_with_matching_saved_rationale(self):
        from pptx import Presentation
        from ai_server.app.proposal_layout_v10 import calculation_pages
        from ai_server.app.proposal_presentation_v4 import PRESENTATION_TEMPLATE_PATH
        report = {'planning_decision': bind_rationale(decision(), pack()),
                  'evidence_sources': pack()['sources'] + pack()['benchmark_cases']}
        prs = Presentation(str(PRESENTATION_TEMPLATE_PATH))
        original_count = len(prs.slides)
        original_xml = [s._element.xml for s in prs.slides]
        calculation_pages(prs, report, original_count-1)
        self.assertEqual(len(prs.slides), original_count+3)
        self.assertEqual([s._element.xml for s in list(prs.slides)[:original_count]], original_xml)
        page = prs.slides[-1]
        page_text = '\n'.join(s.text for s in page.shapes if s.has_text_frame)
        self.assertIn('사실에서 적용 가설까지', page_text)
        self.assertIn('6.6%', page_text)
        self.assertIn('완주 후 시장 방문', page_text)
        self.assertIn('below_baseline_on_test', page.notes_slide.notes_text_frame.text)
        for shape in page.shapes:
            self.assertGreaterEqual(shape.left, 0)
            self.assertLessEqual(shape.left+shape.width, prs.slide_width)
            self.assertLessEqual(shape.top+shape.height, prs.slide_height)
        old = Presentation(str(PRESENTATION_TEMPLATE_PATH))
        calculation_pages(old, {}, len(old.slides)-1)
        self.assertEqual(len(old.slides), original_count+2)

    def test_binding_preserves_exact_facts_hypothesis_status_and_input(self):
        original_pack, original_decision = pack(), decision()
        before = deepcopy((original_pack, original_decision))
        result = bind_rationale(original_decision, original_pack)
        view = report_rationale({'planning_decision': result})
        self.assertEqual(view['facts']['dataset:1:1']['record'], original_pack['sources'][0])
        forecast = next(v for v in view['facts'].values() if v['kind'] == 'forecast')
        self.assertEqual(forecast['record'], original_pack['snapshot']['ml_analysis']['signals'][0])
        self.assertEqual(result['selection_status'], 'needs_evidence')
        self.assertEqual(view['reasoning'], original_decision['design_candidates'][0]['reasoning'])
        self.assertEqual((original_pack, original_decision), before)
        self.assertIn('ml:1:v1', result['design_candidates'][0]['evidence_source_ids'])
        self.assertEqual(len(rationale_sections({'planning_decision': result})), 4)

    def test_fact_ids_stable_when_signal_order_changes(self):
        p = pack()
        second = deepcopy(p['snapshot']['ml_analysis']['signals'][0])
        second['metric'] = 'spending_krw'
        p['snapshot']['ml_analysis']['signals'].append(second)
        before = fact_catalog(p)
        p['snapshot']['ml_analysis']['signals'].reverse()
        self.assertEqual(fact_catalog(p), before)

    def test_unknown_fact_and_changed_case_do_not_become_verified_explanation(self):
        bad = decision()
        bad['design_candidates'][0]['reasoning']['fact_ids'].append('made-up')
        bound = bind_rationale(bad, pack())
        self.assertIsNone(report_rationale({'planning_decision': bound}))
        bound = bind_rationale(decision(), pack())
        bound['design_candidates'][0]['case_source_ids'] = ['case:other']
        self.assertIsNone(report_rationale({'planning_decision': bound}))
        bound = bind_rationale(decision(), pack())
        bound['design_candidates'][0]['measurement_plan'] = '새 설계'
        self.assertIsNone(report_rationale({'planning_decision': bound}))
        bound = bind_rationale(decision(), pack())
        bound['selection_reason'] = '새 이유'
        self.assertIsNone(report_rationale({'planning_decision': bound}))

    def test_old_reports_have_no_invented_trace(self):
        old = decision()
        old['design_candidates'][0].pop('reasoning')
        self.assertEqual(bind_rationale(old, pack()), old)
        self.assertEqual(rationale_sections({'planning_decision': old}), [])

    def test_planner_tool_receives_same_bound_facts_and_hypothesis(self):
        p = pack()
        p['transfer_assessment'] = bind_rationale(decision(), p)
        tool = EvidenceTools({'evidence_pack': p}, task='planner')
        delivered = tool._decision()
        self.assertEqual(delivered['rationale_evidence'], p['transfer_assessment']['rationale_evidence'])
        self.assertEqual(delivered['selected_candidate']['reasoning'], decision()['design_candidates'][0]['reasoning'])


class AgentContractTest(unittest.IsolatedAsyncioTestCase):
    async def test_request_schema_rejects_made_up_ids_and_tool_gets_guide(self):
        router = SimpleNamespace(generate=AsyncMock(return_value=decision()))
        result = await TransferabilityAgent(api_key='', model='', llm_router=router).assess(evidence_pack=pack(), structured_reasoning=True)
        request = router.generate.call_args.args[0]
        reasoning_schema = request.schema['properties']['design_candidates']['items']['properties']['reasoning']
        valid = decision()['design_candidates'][0]['reasoning']
        Draft202012Validator(reasoning_schema).validate(valid)
        invalid = deepcopy(valid)
        invalid['fact_ids'] = ['invented-fact']
        with self.assertRaises(ValidationError):
            Draft202012Validator(reasoning_schema).validate(invalid)
        overview = EvidenceTools(request.input_payload, task='transferability').overview()
        self.assertEqual(set(overview['reasoning_guide']['facts']), set(fact_catalog(pack())))
        self.assertNotIn('record', next(iter(overview['reasoning_guide']['facts'].values())))
        self.assertEqual(overview['reasoning_guide']['originals']['forecast'], 'get_ml_forecast')
        self.assertIsNotNone(report_rationale({'planning_decision': result}))

    async def test_default_generation_does_not_force_experimental_contract(self):
        router = SimpleNamespace(generate=AsyncMock(return_value=decision()))
        await TransferabilityAgent(api_key='', model='', llm_router=router).assess(evidence_pack=pack())
        request = router.generate.call_args.args[0]
        self.assertNotIn('reasoning', request.schema['properties']['design_candidates']['items']['properties'])
        self.assertNotIn('reasoning_guide', request.input_payload)


if __name__ == '__main__':
    unittest.main()
